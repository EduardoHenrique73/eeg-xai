"""Apply verified editorial corrections to the supplied TCC II DOCX.

The input is never overwritten. Moving OOXML paragraph nodes preserves the
embedded figures at their original resolution and their existing captions.
"""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.shared import Pt


SOURCE = Path(r"C:\Users\Foco\Downloads\TCC_II_Eduardo_Marques_Formatado_ABNT.docx")
OUTPUT = Path(__file__).resolve().parents[1] / "docs/tcc/TCC II - Eduardo Marques - ABNT Corrigido.docx"


def replace_text(paragraph, text: str) -> None:
    paragraph.text = text


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--toc-pdf", type=Path)
    args = parser.parse_args()
    document = Document(SOURCE)
    paragraphs = document.paragraphs

    replace_text(
        paragraphs[92],
        "• Dados e métricas de IA: a base final foi a CHB-MIT (SHOEB, 2009), "
        "com registros contínuos multicanais em EDF e anotações temporais das crises. "
        "Foram destinados ao treino 77 EDFs de 17 grupos de pacientes e à calibração "
        "13 EDFs de chb16, chb19, chb20 e chb24. A separação ocorreu por paciente "
        "antes do janelamento. Foram avaliados precisão, sensibilidade e F1 por EDF "
        "e por evento, F1 localizado, IoU temporal e falsos alarmes por hora.",
    )
    replace_text(
        paragraphs[116],
        "7 – Solicitar Análise Preditiva: o sistema deve acionar o modelo CNN-BiLSTM "
        "sobre um exame EDF já importado para gerar probabilidades ictais por intervalo "
        "e destacar trechos suspeitos para revisão médica.",
    )
    replace_text(
        paragraphs[126],
        "Segundo Pressman e Maxim (2016), os casos de uso descrevem as interações "
        "entre atores e funcionalidades do sistema. Neste protótipo, o médico é o ator "
        "principal; a extração de características e a inferência são operações internas "
        "acionadas após a solicitação de análise de um exame previamente importado.",
    )
    table = document.tables[5]
    table.cell(9, 1).text = (
        "2 – Extrair 15 características por canal e por janela, incluindo Entropia de "
        "Shannon e medidas de Dinâmica Simbólica."
    )
    table.cell(10, 1).text = (
        "3 – Alimentar a rede CNN-BiLSTM com as sequências de características para "
        "estimar a probabilidade ictal em cada passo temporal."
    )

    replace_text(
        paragraphs[146],
        "A aplicação adota uma organização cliente-servidor: a interface React se "
        "comunica com a API FastAPI por HTTP, enquanto o servidor mantém as regras de "
        "negócio, a persistência e o processamento de EEG. Essa separação facilita a "
        "manutenção dos componentes (RICHARDS, 2015). A análise de IA é solicitada "
        "separadamente após a importação do EDF e executada em segundo plano.",
    )
    replace_text(
        paragraphs[147],
        "O frontend é uma aplicação React e TypeScript. O traçado multicanal utiliza "
        "Recharts e oferece navegação temporal, seleção de canais e sobreposição de "
        "trechos suspeitos. O backend, em Python e FastAPI, utiliza MNE-Python na "
        "leitura do EDF e TensorFlow/Keras na inferência da CNN-BiLSTM. A instalação "
        "local usa banco relacional configurável; PostgreSQL é previsto no projeto, "
        "mas SQLite também pode ser usado no ambiente de desenvolvimento.",
    )
    replace_text(
        paragraphs[148],
        "Após o envio do EDF, a API valida e armazena o arquivo e registra o exame. "
        "Quando o médico solicita a análise, o pipeline extrai 15 características "
        "por canal em janelas temporais e forma sequências para a CNN-BiLSTM, que "
        "produz probabilidades ictais ao longo do tempo. A explicação Gradient SHAP "
        "é calculada para a sequência de pico e projetada por canal e intervalo de "
        "segundos na interface. Essa projeção não atribui relevância a cada amostra "
        "bruta do sinal e não substitui a interpretação médica.",
    )
    replace_text(paragraphs[155], paragraphs[155].text.replace("utilizador", "usuário"))
    replace_text(
        paragraphs[157],
        "3. Processamento Inteligente Assíncrono: após a importação, o médico solicita "
        "a análise do exame. O servidor executa em segundo plano a extração das 15 "
        "características por canal e a inferência CNN-BiLSTM. As probabilidades são "
        "agregadas em trechos temporais suspeitos; o SHAP explica a sequência de pico. "
        "Os metadados e resultados são registrados no banco relacional configurado, "
        "enquanto EDF e imagens permanecem em disco.",
    )
    replace_text(
        paragraphs[158],
        "4. Conclusão Clínico-Administrativa: a interface apresenta o EEG, os trechos "
        "suspeitos e, quando disponível, a relevância SHAP por canal e intervalo. "
        "O médico pode registrar seu parecer no laudo. A saída da IA é apoio à revisão, "
        "não diagnóstico autônomo.",
    )

    # Keep the original figure paragraphs; only change text around them.
    for index in (172, 173, 174, 175, 176, 180, 184, 188, 192, 196, 200):
        text = paragraphs[index].text
        for old, new in (
            ("ecrãs", "telas"),
            ("utilizador", "usuário"),
            ("auxílio ao diagnóstico neurológico", "apoio à análise neurológica"),
            ("que se seguem", "a seguir"),
            ("foram concebidos", "foram elaborados"),
        ):
            text = text.replace(old, new)
        if text != paragraphs[index].text:
            replace_text(paragraphs[index], text)
    replace_text(
        paragraphs[200],
        "A Figura 11 apresenta a tela de configurações, com dados profissionais, "
        "montagem padrão de canais, limiar de confiança e controle de exibição do SHAP. "
        "Essas preferências de interface não substituem a calibração científica do modelo.",
    )

    # Move the complete implemented-interface section, including image nodes,
    # immediately after the Results heading.
    results_heading = paragraphs[222]._element
    cursor = results_heading
    for paragraph in paragraphs[204:222]:
        element = paragraph._element
        element.getparent().remove(element)
        cursor.addnext(element)
        cursor = element
    replace_text(paragraphs[204], "5.1 Interface implementada")
    replace_text(
        paragraphs[205],
        "As Figuras 12 a 16 documentam a interface implementada no ambiente local: "
        "painel principal, gestão de pacientes, configurações, análise multicanal "
        "com sobreposição SHAP e redação do laudo médico.",
    )
    for index, title in (
        (223, "5.2 Protocolo experimental"),
        (227, "5.3 Desempenho do modelo"),
        (238, "5.4 Evolução experimental e estabilidade entre seeds"),
        (243, "5.5 Análise dos erros e generalização"),
        (245, "5.6 Verificação do software e dos artefatos"),
    ):
        replace_text(paragraphs[index], title)
    replace_text(paragraphs[228], "Quadro 8 – Resultados da v43 na calibração e na avaliação histórica.")

    # One uncited reference points to a dictionary site rather than the article.
    for paragraph in paragraphs:
        if paragraph.text.lstrip().upper().startswith("GROVES, P."):
            paragraph._element.getparent().remove(paragraph._element)
            break
    else:
        raise ValueError("Referência GROVES não encontrada")

    budget = document.tables[9]
    budget.cell(4, 0).text = "TOTAL: R$ 3.650,00"
    for cell in budget.rows[1].cells + budget.rows[2].cells + budget.rows[3].cells:
        cell.text = cell.text.replace("R$100,00", "R$ 100,00")
    appendix = next(p for p in document.paragraphs if p.text.startswith("APÊNDICE A"))
    following = appendix._element.getnext()
    while following is not None and not following.tag.endswith("}p"):
        following = following.getnext()
    if following is not None:
        from docx.text.paragraph import Paragraph
        replace_text(
            Paragraph(following, appendix._parent),
            "O desenvolvimento foi organizado conforme o cronograma do Quadro 10. "
            "Os recursos previstos são apresentados no Quadro 9, cujo total é "
            "R$ 3.650,00.",
        )
    for paragraph in document.paragraphs:
        if paragraph.text.startswith("Quadro 8 – Descrição Orçamentária"):
            replace_text(paragraph, paragraph.text.replace("Quadro 8", "Quadro 9", 1))
        elif paragraph.text.startswith("Quadro 9 – Cronograma"):
            replace_text(paragraph, paragraph.text.replace("Quadro 9", "Quadro 10", 1))

    # This document's summary is a static table, not an updatable Word field.
    summary = document.tables[0]
    interface_row = next(
        row._tr for row in summary.rows
        if row.cells[0].text.strip().startswith("4.9 Interface implementada")
    )
    results_row = next(
        row._tr for row in summary.rows
        if row.cells[0].text.strip() == "5 RESULTADOS"
    )
    interface_row.getparent().remove(interface_row)
    results_row.addnext(interface_row)
    for row in summary.rows:
        label = row.cells[0].text.strip()
        if label.startswith("4.9 Interface implementada"):
            row.cells[0].text = "5.1 Interface implementada"
        elif label.startswith("5.1 Protocolo"):
            row.cells[0].text = "5.2 Protocolo experimental"
        elif label.startswith("5.2 Desempenho"):
            row.cells[0].text = "5.3 Desempenho do modelo"
        elif label.startswith("5.3 Evolução"):
            row.cells[0].text = "5.4 Evolução experimental e estabilidade entre seeds"
        elif label.startswith("5.4 Análise"):
            row.cells[0].text = "5.5 Análise dos erros e generalização"
        elif label.startswith("5.5 Verificação"):
            row.cells[0].text = "5.6 Verificação do software e dos artefatos"

    if args.toc_pdf:
        from pypdf import PdfReader

        pages = [page.extract_text() or "" for page in PdfReader(args.toc_pdf).pages]
        for row in summary.rows:
            label = row.cells[0].text.strip()
            matches = [number for number, page in enumerate(pages[3:], start=4) if label in page]
            if not matches:
                raise ValueError(f"Título não encontrado no PDF: {label}")
            row.cells[1].text = str(matches[0])

    for row in summary.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.line_spacing = 1.0
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(9.5)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
