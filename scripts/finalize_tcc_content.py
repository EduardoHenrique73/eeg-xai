from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


SOURCE = Path("docs/tcc/TCC II - Eduardo Marques - Revisao de Conteudo.docx")
OUTPUT = Path("docs/tcc/TCC II - Eduardo Marques - Versao Final Corrigida.docx")


def replace_paragraph_text(paragraph: Paragraph, old: str, new: str) -> None:
    full_text = paragraph.text
    if old not in full_text:
        raise RuntimeError(f"Texto não encontrado: {old}")
    updated = full_text.replace(old, new)
    if paragraph.runs:
        paragraph.runs[0].text = updated
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(updated)


def find_paragraph(document: Document, exact_text: str) -> Paragraph:
    for paragraph in document.paragraphs:
        if paragraph.text.strip() == exact_text:
            return paragraph
    raise RuntimeError(f"Parágrafo não encontrado: {exact_text}")


def insert_before(target: Paragraph, text: str, reference: Paragraph) -> Paragraph:
    element = OxmlElement("w:p")
    if reference._p.pPr is not None:
        element.append(deepcopy(reference._p.pPr))
    target._p.addprevious(element)
    paragraph = Paragraph(element, target._parent)
    paragraph.add_run(text)
    return paragraph


def insert_after(target: Paragraph, text: str, reference: Paragraph) -> Paragraph:
    element = OxmlElement("w:p")
    if reference._p.pPr is not None:
        element.append(deepcopy(reference._p.pPr))
    target._p.addnext(element)
    paragraph = Paragraph(element, target._parent)
    paragraph.add_run(text)
    return paragraph


def remove_paragraph(paragraph: Paragraph) -> None:
    paragraph._element.getparent().remove(paragraph._element)


def main() -> None:
    document = Document(SOURCE)

    replacements = {
        "usuárioes": "usuários",
        "A conceção destas interfaces": "A concepção dessas interfaces",
        "os principais telas projetados": "as principais telas projetadas",
        "A secção lateral esquerda": "A seção lateral esquerda",
        "seleção dinâmica de elétrodos": "seleção dinâmica de eletrodos",
        "modelo preditivo (CNN-LSTM)": "modelo preditivo (CNN-BiLSTM)",
        "num ambiente limpo e responsivo": "em um ambiente limpo e responsivo",
        "até à área central": "até a área central",
    }
    for paragraph in document.paragraphs:
        for old, new in replacements.items():
            if old in paragraph.text:
                replace_paragraph_text(paragraph, old, new)

    item_processamento = find_paragraph(
        document,
        "3. Processamento Inteligente Assíncrono: após a importação, o médico solicita a análise do exame. "
        "O servidor executa em segundo plano a extração das 15 características por canal e a inferência "
        "CNN-BiLSTM. As probabilidades são agregadas em trechos temporais suspeitos; o SHAP explica a "
        "sequência de pico. Os metadados e resultados são registrados no banco relacional configurado, "
        "enquanto EDF e imagens permanecem em disco.",
    )
    body_reference = find_paragraph(
        document,
        "As Figuras 12 a 16 documentam a interface implementada no ambiente local: painel principal, "
        "gestão de pacientes, configurações, análise multicanal com sobreposição SHAP e redação do laudo médico.",
    )
    insert_after(
        item_processamento,
        "Na Figura 3, os blocos roxos representam o pipeline conceitual definido no planejamento inicial. "
        "Na implementação final, essas etapas foram consolidadas na extração de 15 características "
        "tempo-frequência por canal e na inferência temporal pela rede CNN-BiLSTM.",
        body_reference,
    )

    descriptions = {
        "Figura 12 – Painel principal da versão implementada.": (
            "A Figura 12 apresenta o painel principal da versão implementada após a autenticação. A tela "
            "reúne os indicadores de pacientes cadastrados, exames em análise e laudos emitidos, além de "
            "atalhos para as atividades mais frequentes do fluxo clínico."
        ),
        "Figura 13 – Gestão de pacientes da versão implementada.": (
            "A Figura 13 apresenta a gestão de pacientes, com pesquisa por nome ou CPF, cadastro de novos "
            "registros e ações para edição, exclusão e abertura de um novo exame vinculado ao prontuário."
        ),
        "Figura 14 – Configurações do perfil, montagem, limiar e SHAP.": (
            "A Figura 14 mostra as configurações do perfil médico e das preferências de análise. Nessa tela, "
            "o usuário pode definir a montagem padrão de canais, ajustar o limiar de exibição e habilitar ou "
            "desabilitar a explicação SHAP."
        ),
        "Figura 15 – Visualização multicanal com trecho suspeito e SHAP sobreposto por canal e tempo.": (
            "A Figura 15 apresenta a área central de revisão do EEG, com navegação temporal, seleção de "
            "canais, candidatos suspeitos e sobreposição da relevância Gradient SHAP por canal e intervalo. "
            "A faixa laranja delimita o trecho suspeito, enquanto a intensidade vermelha representa "
            "contribuição positiva para a saída do modelo. A projeção está associada às janelas de "
            "características da sequência de pico e não corresponde a uma atribuição para cada amostra bruta."
        ),
        "Figura 16 – Área de emissão do parecer médico após a análise computacional.": (
            "A Figura 16 apresenta a área de emissão do parecer médico. O especialista pode revisar o "
            "traçado e os resultados da IA, registrar sua interpretação clínica e concluir o laudo, que "
            "permanece associado ao exame no banco de dados."
        ),
    }
    for caption, description in descriptions.items():
        insert_before(find_paragraph(document, caption), description, body_reference)

    old_figure_15_note = find_paragraph(
        document,
        "Na Figura 15, a faixa laranja delimita o trecho suspeito e a intensidade vermelha representa "
        "contribuição positiva do Gradient SHAP. A projeção está associada às janelas de features da "
        "sequência de pico; portanto, não corresponde a uma atribuição para cada amostra bruta do sinal.",
    )
    remove_paragraph(old_figure_15_note)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT.resolve())


if __name__ == "__main__":
    main()
