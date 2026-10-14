"""Apply narrowly scoped factual corrections to the formatted TCC II."""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/tcc/TCC II - Eduardo Marques - ABNT Corrigido.docx"
OUTPUT = ROOT / "docs/tcc/TCC II - Eduardo Marques - Revisao de Conteudo.docx"


def change(paragraph, prefix: str, value: str) -> None:
    if not paragraph.text.startswith(prefix):
        raise ValueError(f"Trecho inesperado: {paragraph.text[:100]!r}")
    paragraph.text = value


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--toc-pdf", type=Path)
    args = parser.parse_args()
    doc = Document(SOURCE)
    p = doc.paragraphs

    change(p[30], "• Extrair características", "• Investigar Entropia de Shannon e Dinâmica Simbólica como características candidatas e selecionar, por avaliação experimental, a representação de entrada da CNN-BiLSTM temporal.")
    change(p[53], "Na persistência,", "Na versão final, o PostgreSQL armazena os dados cadastrais e metadados dos exames, enquanto os arquivos EDF e os artefatos de IA permanecem no sistema de arquivos, com seus caminhos referenciados pela aplicação. A conexão e a leitura dos registros foram verificadas no fluxo local da API. Não foram utilizados pgvector nem banco vetorial.")
    change(p[56], "A abordagem proposta", "No TCC I, a proposta de CNN-LSTM partia da análise temporal do EEG e previa Entropia de Shannon e Dinâmica Simbólica como entradas possíveis. Durante a implementação, diferentes representações foram avaliadas. A versão final utiliza uma CNN-BiLSTM sobre sequências de características tempo-frequência extraídas por canal, e não sobre todas as amostras brutas da onda. A parte convolucional aprende padrões nessa representação; a parte recorrente modela sua evolução temporal. Assim, a arquitetura implementada mantém o objetivo de detectar intervalos suspeitos, mas não corresponde literalmente à entrada bruta descrita na proposta inicial (HOCHREITER; SCHMIDHUBER, 1997; YAMASHITA et al., 2018).")
    change(p[57], "A escolha de uma abordagem", "O trabalho anterior de Moreira et al. (2025) integra o contexto acadêmico que motivou a investigação de características matemáticas e modelos de IA. Seus resultados numéricos não são utilizados aqui como evidência de desempenho em EEG, pois não foi possível conferir no relatório original o protocolo e os valores atribuídos a essa comparação. A escolha da CNN-BiLSTM e de sua representação de entrada fundamenta-se nos experimentos controlados descritos neste trabalho, sem transferir acurácias de outro estudo para a avaliação da CHB-MIT.")
    change(p[81], "O backend foi", "O backend foi desenvolvido com Python 3.11.9, FastAPI 0.115.6 e SQLAlchemy 2.0.36. Na configuração final, o PostgreSQL realiza a persistência relacional; EDFs, modelos e mapas permanecem no sistema de arquivos. A conexão com a API e a leitura dos registros foram verificadas após a migração dos dados locais. O ambiente Python é isolado por venv.")
    change(p[87], "• Incremento II", "• Incremento II (Núcleo de Inteligência): o extrator histórico implementou Entropia de Shannon e medidas de Dinâmica Simbólica, que foram investigadas como alternativas de representação. Após os experimentos, a CNN-BiLSTM final passou a receber 15 características tempo-frequência por canal, incluindo entropia espectral, mas não aquelas características históricas. A rede produz probabilidades ictais por passo temporal. A separação dos grupos ocorreu antes do janelamento e da amostragem.")
    change(p[147], "O frontend é", "O frontend é uma aplicação React e TypeScript. O traçado multicanal utiliza Recharts e oferece navegação temporal, seleção de canais e sobreposição de trechos suspeitos. O backend, em Python e FastAPI, utiliza MNE-Python na leitura do EDF e TensorFlow/Keras na inferência da CNN-BiLSTM. Os metadados são persistidos no PostgreSQL configurado para a versão final.")
    change(p[152], "• Camada de Servidor", "• Camada de Servidor (em verde): representa o backend FastAPI, responsável pela lógica de negócio, validação dos arquivos e persistência dos metadados no PostgreSQL.")
    change(p[224], "O modelo final selecionado", "O modelo final selecionado foi a versão v43 com seed 42. A arquitetura CNN-BiLSTM utiliza fusão tardia de 23 canais e 15 características tempo-frequência por canal, entre elas entropia espectral. Não utiliza a feature histórica de Entropia de Shannon nem medidas de Dinâmica Simbólica. As janelas têm 4 segundos, passo de 2 segundos e sequências de 8 passos. O treinamento utilizou 77 EDFs de 17 grupos identificados na base; a calibração reuniu 13 EDFs de chb16, chb19, chb20 e chb24, separados antes do janelamento.")
    change(doc.tables[5].cell(9, 1).paragraphs[0], "2 – Extrair", "2 – Extrair 15 características tempo-frequência por canal e por janela, incluindo entropia espectral, para formar sequências temporais.")
    for table_index in (1, 2):
        cell = doc.tables[table_index].cell(11, 1)
        if "PostgreSQL" not in cell.text:
            raise ValueError(f"PostgreSQL não encontrado no quadro {table_index}")

    if not p[246].text.startswith("A integração foi verificada"):
        raise ValueError("Trecho de verificação do software inesperado")
    p[246].text += (
        " Na configuração final, a API iniciou com PostgreSQL e recuperou os "
        "registros migrados de usuários, pacientes, exames e predições. A comparação "
        "campo a campo com a origem confirmou a integridade dos registros."
    )

    anchor = p[239]._element
    additions = [
        "A proposta registrada no TCC I previa a combinação das bases Bonn e CHB-MIT, validação k-fold e uso de Entropia de Shannon e Dinâmica Simbólica em uma CNN-LSTM. Um ensaio inicial de k-fold por janelas foi exploratório e não constitui a validação do modelo final. A avaliação passou a separar grupos antes do janelamento para examinar generalização interpaciente e evitar compartilhamento de janelas do mesmo grupo entre etapas. A base Bonn foi retirada da avaliação temporal final pelas diferenças de estrutura e anotação em relação aos EDFs contínuos da CHB-MIT.",
        "A seleção das 15 características espectrais por canal decorreu das comparações de representação documentadas nos experimentos. A entropia espectral da v43 não equivale à feature histórica de Entropia de Shannon calculada diretamente do sinal, e a Dinâmica Simbólica não alimenta o modelo final. Também foi testada uma representação de sinal bruto, que não substituiu a escolhida. Essas alterações são decisões tomadas ao longo do desenvolvimento, não requisitos definidos dessa forma desde o início. Os resultados abaixo referem-se exclusivamente ao protocolo e aos artefatos da versão final.",
    ]
    for value in additions:
        element = OxmlElement("w:p")
        anchor.addprevious(element)
        new = Paragraph(element, p[239]._parent)
        new.style = p[239].style
        new.paragraph_format.first_line_indent = p[239].paragraph_format.first_line_indent
        new.text = value

    if args.toc_pdf:
        from pypdf import PdfReader

        pages = [page.extract_text() or "" for page in PdfReader(args.toc_pdf).pages]
        for row in doc.tables[0].rows:
            label = row.cells[0].text.strip()
            matches = [i for i, page in enumerate(pages[4:], start=5) if label in page]
            if not matches:
                raise ValueError(f"Título não encontrado no PDF: {label}")
            row.cells[1].text = str(matches[0])

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
