"""Repair the TCC II document while retaining the original institutional layout."""

from __future__ import annotations

import re
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/tcc/TCC II - Eduardo Marques - Versao Final.docx"
OUTPUT = ROOT / "docs/tcc/TCC II - Eduardo Marques - Versao Final Revisada.docx"
TOC_PAGES = ROOT / "docs/tcc/render_revisao/toc_pages.json"


def replace_text(paragraph, value):
    paragraph.clear()
    paragraph.add_run(value)


def before(doc, anchor, value, style="normal"):
    paragraph = doc.add_paragraph(value, style=style)
    anchor._p.addprevious(paragraph._p)
    return paragraph


def main():
    doc = Document(SOURCE)
    paragraphs = doc.paragraphs

    # The previous builder retained an empty Word TOC content control over a
    # manually typed contents paragraph. Remove that control before rebuilding.
    for sdt in list(doc.element.body.xpath(".//w:sdt")):
        if sdt.xpath('.//w:instrText[contains(text(), "TOC")]'):
            sdt.getparent().remove(sdt)

    replace_text(paragraphs[63], "")
    paragraphs[62].text = "SUMÁRIO"
    paragraphs[62].alignment = WD_ALIGN_PARAGRAPH.CENTER

    title = (
        "DESENVOLVIMENTO DE PLATAFORMA WEB PARA APOIO À ANÁLISE DE EEG "
        "COM INTELIGÊNCIA ARTIFICIAL EXPLICÁVEL"
    )
    for index in (19, 40):
        replace_text(paragraphs[index], title)
        paragraphs[index].alignment = WD_ALIGN_PARAGRAPH.CENTER

    replace_text(
        paragraphs[68],
        paragraphs[68].text.replace(
            "o presente trabalho propõe o desenvolvimento de um protótipo funcional",
            "o presente trabalho desenvolveu um protótipo funcional",
        ),
    )

    # Remove a dangling URL introduced by the previous reference insertion.
    replace_text(
        paragraphs[507],
        "META PLATFORMS. React: the library for web and native user interfaces. "
        "Disponível em: https://react.dev/. Acesso em: 13 out. 2026.",
    )
    replace_text(
        paragraphs[514],
        "REISMAN, M. EHRs: the challenge of making electronic data usable and interoperable. "
        "Pharmacy and Therapeutics, v. 42, n. 9, p. 572-575, 2017. "
        "Disponível em: https://pmc.ncbi.nlm.nih.gov/articles/PMC5565131/. "
        "Acesso em: 13 out. 2026.",
    )

    # Continue the figure sequence already used by the original TCC.
    captions = [
        ("Painel principal da versão implementada", 12),
        ("Gestão de pacientes da versão implementada", 13),
        ("Configurações do perfil, montagem", 14),
        ("Visualização multicanal com trecho suspeito", 15),
        ("Área de emissão do parecer", 16),
        ("Comparação das métricas da v43", 17),
        ("Eventos localizados e falsos alarmes", 18),
        ("Verificação de integridade dos artefatos", 19),
    ]
    for paragraph in doc.paragraphs:
        for fragment, number in captions:
            if fragment in paragraph.text and paragraph.text.startswith("Figura "):
                replace_text(paragraph, re.sub(r"^Figura \d+", f"Figura {number}", paragraph.text))
    for paragraph in doc.paragraphs:
        if "Na Figura 14, a faixa" in paragraph.text:
            replace_text(paragraph, paragraph.text.replace("Figura 14", "Figura 15"))
        if "Figura 2" in paragraph.text and "metodologia" in paragraph.text.lower():
            replace_text(paragraph, paragraph.text.replace("Figura 2", "Figura 1"))
        if "As telas finais substituem" in paragraph.text:
            replace_text(
                paragraph,
                "As figuras seguintes documentam a interface implementada no ambiente local. "
                "O fluxo contempla painel, pacientes, configurações, inspeção multicanal do EEG "
                "com relevância SHAP e elaboração do parecer médico.",
            )

    # Insert missing interpretation of the experiments in the results chapter.
    anchor = next(p for p in doc.paragraphs if p.text.startswith("5.3 Verificação"))
    additions = [
        ("5.3 Evolução experimental e estabilidade entre seeds", "Heading 2"),
        (
            "A seleção da v43 não decorreu de uma única execução. Nas cinco seeds 42 a 46, "
            "a calibração apresentou média de 13,8 eventos detectados em 20, F1 por evento "
            "de 65,7%, F1 localizado de 81,9%, F1 por EDF de 96,4% e 0,64 falso alarme por hora. "
            "A seed 42 alcançou 15/20 eventos, F1 localizado de 90,9% e 0,55 falso alarme "
            "por hora. A variação entre seeds impede tratar esse melhor resultado isolado como "
            "desempenho esperado em novos pacientes.",
            "normal",
        ),
        (
            "Em comparação controlada com a CNN-BiLSTM anterior nas seeds 42 a 44, "
            "a fusão tardia da v43 elevou a média de eventos de 13,33 para 14,00 em 20, "
            "o F1 localizado de 75,0% para 83,4% e o F1 por EDF de 92,9% para 97,1%. "
            "Os falsos alarmes passaram de 0,58 para 0,63 por hora. Esses números são de "
            "calibração no mesmo protocolo e não demonstram generalização externa.",
            "normal",
        ),
        (
            "Experimentos posteriores avaliaram multirresolução, seleção de checkpoint e "
            "expansão dirigida dos dados. A v53, com 26 EDF no treino e três seeds, obteve "
            "média de 12,33/20 eventos, F1 localizado de 77,9%, F1 por EDF de 92,5% "
            "e 0,602 falso alarme por hora. Embora algumas crises tenham melhorado, o conjunto "
            "não superou a referência v43 no mesmo protocolo. Por isso, a v53 não substituiu "
            "o modelo integrado à aplicação.",
            "normal",
        ),
        (
            "A mineração de negativos difíceis foi examinada nas versões v55 e v56, usando "
            "somente pacientes de treino. A configuração forte v55 reduziu a média para "
            "12,67/20 eventos e F1 localizado de 71,3%, sem redução relevante dos falsos "
            "alarmes. A configuração leve v56 manteve 15/20 eventos na seed 42, mas reduziu "
            "o F1 localizado de 90,9% para 85,7%, com os mesmos 0,55 falso alarme por hora. "
            "Ambas foram rejeitadas para implantação.",
            "normal",
        ),
        ("5.4 Análise dos erros e generalização", "Heading 2"),
        (
            "A avaliação histórica em chb09, chb15 e chb18 localizou 4 de 9 eventos, "
            "com F1 localizado de 46,2% e 2,08 falsos alarmes por hora. Esses pacientes "
            "haviam sido consultados durante o desenvolvimento e, portanto, os resultados "
            "descrevem a dificuldade interpaciente observada, sem constituir teste externo "
            "intocado. Crises não localizadas podem refletir baixa evidência ictal, "
            "fragmentação temporal ou escolha de candidato principal incorreto. "
            "A interpretação por evento é necessária porque F1 por EDF não identifica "
            "qual desses erros ocorreu.",
            "normal",
        ),
    ]
    for value, style in additions:
        before(doc, anchor, value, style)
    replace_text(anchor, "5.5 Verificação do software e dos artefatos")
    anchor.style = "Heading 2"

    protocol = next(p for p in doc.paragraphs if p.text.startswith("5.2 Desempenho"))
    before(
        doc,
        protocol,
        "A base Bonn, prevista no planejamento do TCC I, não foi combinada à avaliação "
        "final. Seus segmentos curtos e previamente selecionados não preservam o mesmo "
        "protocolo de registros contínuos multicanais com início e término de crises do "
        "CHB-MIT. A mistura direta dificultaria interpretar métricas de localização "
        "temporal e falsos alarmes por hora. Assim, os resultados apresentados nesta "
        "pesquisa referem-se ao CHB-MIT; Bonn permanece como referência do planejamento.",
    )

    comparison = next(p for p in doc.paragraphs if p.text.startswith("5.3 Evolução"))
    before(
        doc,
        comparison,
        "O ponto operacional da v43 foi escolhido na curva de compromisso entre "
        "sensibilidade por evento e falsos alarmes por hora (FROC), com teto de "
        "0,75 falso alarme por hora na calibração. Os valores apresentados no quadro "
        "representam esse ponto selecionado; não constituem uma curva FROC de teste "
        "externo. A precisão por EDF e a localização são analisadas separadamente, pois "
        "uma classificação correta do arquivo pode coexistir com uma crise perdida.",
    )

    # Resources and the original project schedule are historical material.
    for paragraph in doc.paragraphs:
        if paragraph.text == "8 RECURSOS E CRONOGRAMA":
            replace_text(paragraph, "APÊNDICE A - RECURSOS E CRONOGRAMA")
        elif paragraph.text == "8.1 Recursos":
            replace_text(paragraph, "A.1 Recursos")
        elif paragraph.text == "8.2 Cronograma":
            replace_text(paragraph, "A.2 Cronograma")
        elif paragraph.text == "9 REFERÊNCIAS":
            replace_text(paragraph, "REFERÊNCIAS")

    appendix = next(p for p in doc.paragraphs if p.text == "APÊNDICE A - RECURSOS E CRONOGRAMA")
    references = next(p for p in doc.paragraphs if p.text == "REFERÊNCIAS")
    body = doc.element.body
    start, stop = body.index(appendix._p), body.index(references._p)
    for element in list(body)[start:stop]:
        body.remove(element)
        body.insert(len(body) - 1, element)

    for index in (32, 62):
        paragraphs[index].paragraph_format.page_break_before = True
    for paragraph in doc.paragraphs:
        if paragraph.style.name == "Heading 1" and paragraph.text.strip():
            paragraph.paragraph_format.page_break_before = True

    normal = doc.styles["normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    for section in doc.sections:
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(3)
        section.left_margin = Cm(3)
        section.bottom_margin = Cm(2)
        section.right_margin = Cm(2)

    all_paragraphs = doc.paragraphs
    reference_index = next(i for i, p in enumerate(all_paragraphs) if p.text == "REFERÊNCIAS")
    cover_nodes = {id(p._p) for p in paragraphs[:65]}
    appendix_index = next(i for i, p in enumerate(all_paragraphs) if p.text == "APÊNDICE A - RECURSOS E CRONOGRAMA")
    reference_nodes = {id(p._p) for p in all_paragraphs[reference_index + 1:appendix_index]}
    for paragraph in all_paragraphs:
        if id(paragraph._p) in reference_nodes:
            normalized = re.sub(r"(Disponível em:)\s*(https?://)", r"\1 \2", paragraph.text)
            if normalized != paragraph.text:
                replace_text(paragraph, normalized)
        value = paragraph.text.strip()
        if not value:
            continue
        fmt = paragraph.paragraph_format
        if paragraph.style.name.startswith("Heading"):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            fmt.first_line_indent = Cm(0)
            fmt.line_spacing = 1.0
            fmt.space_before = Pt(12)
            fmt.space_after = Pt(6)
            fmt.keep_with_next = True
            for run in paragraph.runs:
                run.font.name = "Arial"
                run.font.size = Pt(12)
                run.font.bold = True
        elif value.startswith(("Figura ", "Quadro ", "Fonte:")):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            fmt.first_line_indent = Cm(0)
            fmt.line_spacing = 1.0
            fmt.space_after = Pt(3)
            fmt.keep_with_next = value.startswith(("Figura ", "Quadro "))
            for run in paragraph.runs:
                run.font.name = "Arial"
                run.font.size = Pt(10)
        elif value == "SUMÁRIO":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            fmt.line_spacing = 1.0
        elif id(paragraph._p) in cover_nodes:
            continue
        elif id(paragraph._p) in reference_nodes:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            fmt.first_line_indent = Cm(0)
            fmt.line_spacing = 1.0
            fmt.space_after = Pt(6)
        else:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            fmt.first_line_indent = Cm(1.25)
            fmt.line_spacing = 1.5
            fmt.space_after = Pt(0)

    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.line_spacing = 1.0
                    paragraph.paragraph_format.first_line_indent = Cm(0)
                    for run in paragraph.runs:
                        run.font.name = "Arial"
                        run.font.size = Pt(10)

    # A borderless two-column table keeps page numbers aligned in both DOCX and PDF.
    toc_pages = json.loads(TOC_PAGES.read_text(encoding="utf-8")) if TOC_PAGES.exists() else {}
    toc_anchor = paragraphs[63]
    toc_table = doc.add_table(rows=0, cols=2)
    toc_table.autofit = False
    toc_table.columns[0].width = Cm(14.8)
    toc_table.columns[1].width = Cm(1.2)
    for paragraph in [p for p in doc.paragraphs if p.style.name.startswith("Heading") and p.text.strip()]:
        if paragraph.text == "SUMÁRIO":
            continue
        cells = toc_table.add_row().cells
        cells[0].width = Cm(14.8)
        cells[1].width = Cm(1.2)
        cells[0].text = paragraph.text.strip()
        cells[1].text = str(toc_pages.get(paragraph.text.strip(), ""))
        level = int(paragraph.style.name[-1]) if paragraph.style.name[-1].isdigit() else 1
        for column, cell in enumerate(cells):
            item = cell.paragraphs[0]
            item.alignment = WD_ALIGN_PARAGRAPH.RIGHT if column else WD_ALIGN_PARAGRAPH.LEFT
            item.paragraph_format.left_indent = Cm(0.45 * (level - 1)) if column == 0 else Cm(0)
            item.paragraph_format.line_spacing = 1.0
            item.paragraph_format.space_after = Pt(0)
            item.paragraph_format.first_line_indent = Cm(0)
            for run in item.runs:
                run.font.name = "Arial"
                run.font.size = Pt(10)
                run.font.bold = level == 1
    toc_anchor._p.addprevious(toc_table._tbl)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
