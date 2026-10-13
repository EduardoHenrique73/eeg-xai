from __future__ import annotations

from pathlib import Path
import re

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt


SOURCE = Path(r"C:\Users\Foco\Downloads\TCC I - Eduardo Marques.docx")
OUTPUT_DIR = Path(r"C:\Users\Foco\Documents\DEV\eeg-xai\docs\tcc")
OUTPUT = OUTPUT_DIR / "TCC I - Eduardo Marques - Revisado.docx"
DIAGRAM = OUTPUT_DIR / "assets" / "fases_metodologicas.png"


def set_text(paragraph, text: str) -> None:
    """Replace paragraph text while retaining paragraph-level formatting."""
    run_props = None
    if paragraph.runs:
        run_props = paragraph.runs[0]._r.rPr
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)
    run = paragraph.add_run(text)
    if run_props is not None:
        run._r.insert(0, run_props)


def insert_after(paragraph, text: str = ""):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_paragraph = paragraph._parent.add_paragraph()
    new_paragraph._p.getparent().remove(new_paragraph._p)
    new_p.addnext(new_paragraph._p)
    new_paragraph.style = paragraph.style
    if text:
        new_paragraph.add_run(text)
    return new_paragraph


def insert_before(paragraph, text: str = ""):
    new_paragraph = paragraph.insert_paragraph_before(text)
    new_paragraph.style = paragraph.style
    return new_paragraph


def remove_numbering(paragraph) -> None:
    p_pr = paragraph._p.pPr
    if p_pr is None:
        return
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is not None:
        p_pr.remove(num_pr)


def make_method_diagram(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    width, height = 2200, 520
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    try:
        font = ImageFont.truetype("arial.ttf", 36)
        small = ImageFont.truetype("arial.ttf", 31)
    except OSError:
        font = ImageFont.load_default()
        small = font

    labels = [
        "Requisitos",
        "Preparação\ndos dados",
        "Treinamento",
        "Calibração",
        "Avaliação",
        "Integração\ne verificação",
    ]
    margin, gap = 55, 40
    box_width = (width - 2 * margin - gap * (len(labels) - 1)) // len(labels)
    top, bottom = 145, 375
    for index, label in enumerate(labels):
        left = margin + index * (box_width + gap)
        right = left + box_width
        draw.rounded_rectangle(
            (left, top, right, bottom),
            radius=20,
            fill="#EAF3FA",
            outline="#185A8D",
            width=5,
        )
        bbox = draw.multiline_textbbox((0, 0), label, font=small, spacing=8, align="center")
        tx = left + (box_width - (bbox[2] - bbox[0])) / 2
        ty = top + (bottom - top - (bbox[3] - bbox[1])) / 2
        draw.multiline_text((tx, ty), label, font=small, fill="#123A57", spacing=8, align="center")
        if index < len(labels) - 1:
            x1, x2, y = right + 7, right + gap - 7, (top + bottom) // 2
            draw.line((x1, y, x2, y), fill="#185A8D", width=6)
            draw.polygon([(x2, y), (x2 - 18, y - 12), (x2 - 18, y + 12)], fill="#185A8D")
    title = "Etapas executadas no desenvolvimento e na avaliação"
    bbox = draw.textbbox((0, 0), title, font=font)
    draw.text(((width - (bbox[2] - bbox[0])) / 2, 45), title, font=font, fill="#111111")
    image.save(path, dpi=(300, 300))


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    make_method_diagram(DIAGRAM)
    doc = Document(SOURCE)
    p = doc.paragraphs

    replacements = {
        30: "URUAÇU/GO",
        60: "URUAÇU/GO",
        66: (
            "O Eletroencefalograma (EEG) constitui uma ferramenta fundamental para o diagnóstico e o "
            "monitoramento da epilepsia, pois registra com elevada resolução temporal a atividade elétrica "
            "cerebral. A análise de registros extensos, contudo, permanece predominantemente visual, demanda "
            "tempo especializado e está sujeita à fadiga e à variabilidade entre avaliadores (TATUM et al., "
            "2018; ROY et al., 2019). Esse cenário reforça a necessidade de ferramentas computacionais capazes "
            "de apoiar a triagem sem substituir a decisão médica."
        ),
        67: (
            "Embora técnicas de Deep Learning apresentem potencial para reconhecer padrões complexos em EEG, "
            "sua incorporação ao fluxo clínico enfrenta barreiras técnicas e operacionais. Entre elas estão a "
            "fragmentação entre ferramentas de armazenamento e visualização, a dependência de soluções "
            "proprietárias, a dificuldade de processar exames multicanais extensos e a opacidade das redes "
            "neurais profundas. Um classificador isolado não resolve integralmente o problema: o especialista "
            "também precisa importar o exame, navegar pelos canais, localizar os eventos suspeitos, avaliar a "
            "explicação do modelo e registrar seu próprio parecer."
        ),
        68: (
            "Diante dessa lacuna, o presente trabalho propõe o desenvolvimento de um protótipo funcional de "
            "plataforma web para apoio à triagem e à análise de EEG. A solução centraliza a ingestão de arquivos "
            "EDF, a visualização multicanal, a inferência temporal por uma rede CNN-BiLSTM e a apresentação de "
            "trechos suspeitos. A arquitetura cliente-servidor mediada por API REST separa a interface do núcleo "
            "de processamento e permite evoluir cada componente de forma independente."
        ),
        69: (
            "A explicabilidade é incorporada por meio do método SHAP, que relaciona a decisão da rede às "
            "características, aos canais e aos intervalos temporais analisados. O sistema deve ser compreendido "
            "como ferramenta experimental de segunda opinião, com potencial para apoiar a produtividade e "
            "direcionar a revisão do especialista, e não como substituto do diagnóstico clínico. Por lidar com "
            "informações de saúde, sua evolução também deve considerar acessibilidade, privacidade, segurança e "
            "o enquadramento regulatório aplicável a Software como Dispositivo Médico."
        ),
        72: (
            "Desenvolver e avaliar um protótipo funcional de plataforma web para apoio à triagem e à análise de "
            "exames de Eletroencefalograma (EEG), integrando uma arquitetura híbrida CNN-BiLSTM e técnicas de "
            "Inteligência Artificial Explicável (XAI) para localizar temporalmente eventos suspeitos e fornecer "
            "informações de apoio à revisão do especialista."
        ),
        75: "Levantar e documentar requisitos funcionais e não funcionais relacionados à ingestão, ao processamento, à visualização, à segurança e à interpretabilidade de sinais de EEG.",
        76: "Implementar uma interface web responsiva, desenvolvida em React, para importação de EDF, navegação temporal, zoom e seleção dinâmica de canais.",
        77: "Desenvolver um módulo de ingestão capaz de validar arquivos EDF, extrair seus metadados e preservar a identidade dos canais durante a análise.",
        78: "Estruturar o janelamento temporal depois da separação dos pacientes entre treino e calibração, ampliando a quantidade de exemplos sem provocar vazamento de dados.",
        79: "Extrair características tempo-frequência por canal, Entropia de Shannon e medidas de Dinâmica Simbólica, e treinar uma rede CNN-BiLSTM com probabilidade ictal por passo temporal.",
        80: "Integrar o método SHAP para representar a relevância das características por canal e por intervalo temporal sobre o traçado do EEG.",
        81: "Avaliar o modelo em registros públicos da CHB-MIT por EDF, evento e localização temporal, utilizando precisão, sensibilidade, F1, IoU e falsos alarmes por hora.",
        83: (
            "A leitura manual de EEG é extensa e depende de conhecimento especializado, enquanto classificadores "
            "isolados não oferecem necessariamente o fluxo de ingestão, visualização, explicação e registro do "
            "parecer. O projeto busca enfrentar essa lacuna por meio de uma plataforma integrada, preservando a "
            "supervisão humana e apresentando a saída da IA como trecho suspeito, e não como diagnóstico final."
        ),
        84: (
            "A proposta fundamenta-se na experiência adquirida no projeto de iniciação tecnológica "
            "Identificação de Padrões Genéticos para Diagnóstico de Doenças Utilizando Inteligência Artificial, "
            "posteriormente direcionado à análise de sinais bioelétricos (MOREIRA et al., 2025). O trabalho "
            "anterior demonstrou a viabilidade de extrair características matemáticas de séries temporais; o "
            "presente projeto amplia essa contribuição ao incorporar tais técnicas a um fluxo web completo."
        ),
        85: (
            "A disponibilização por navegador reduz a dependência de instalações proprietárias em cada estação "
            "de trabalho e permite centralizar o processamento computacional no servidor. Essa estratégia não "
            "elimina custos de infraestrutura, mas favorece manutenção, atualização e acesso controlado. A "
            "integração entre visualização, modelo temporal e XAI também contribui para que a predição seja "
            "inspecionada no contexto do sinal original."
        ),
        86: (
            "A relevância acadêmica decorre da aplicação conjunta de Engenharia de Software, processamento de "
            "sinais, aprendizado profundo e interpretabilidade. O resultado esperado é um protótipo verificável, "
            "acompanhado de métricas e limitações explícitas, que possa servir de base para estudos posteriores "
            "de generalização interpaciente, usabilidade e validação com especialistas."
        ),
        91: (
            "O Eletroencefalograma (EEG) é reconhecido como recurso fundamental para o diagnóstico e o "
            "monitoramento de epilepsia e outras desordens cerebrais, caracterizando-se como exame não invasivo, "
            "de baixo custo e elevada resolução temporal (NIEDERMEYER; LOPES DA SILVA, 2005; CRAIK et al., "
            "2019). A duração dos registros e a quantidade de canais produzem grande volume de dados, o que "
            "aumenta a carga de revisão dos especialistas."
        ),
        94: "Aumenta a Sensibilidade: redes profundas podem identificar relações não lineares e padrões distribuídos no sinal que são difíceis de representar por regras fixas; o desempenho, contudo, depende da base, do protocolo de validação e da população avaliada (ACHARYA et al., 2018; ROY et al., 2019).",
        97: "Portanto, uma plataforma web de apoio não busca substituir o médico, mas organizar o fluxo de triagem e apresentar evidências computacionais para revisão. A utilidade clínica efetiva depende de validação externa, avaliação com especialistas e controle dos falsos alarmes.",
        99: (
            "O desenvolvimento adotou uma abordagem incremental adaptada a um projeto individual de Pesquisa e "
            "Desenvolvimento. Em vez de reproduzir integralmente papéis coletivos de Scrum ou XP, o trabalho foi "
            "organizado em ciclos de implementação, experimento, avaliação e revisão com a orientação acadêmica. "
            "O Git foi utilizado para versionar o código e manter a rastreabilidade das alterações."
        ),
        100: "Os ciclos curtos permitiram comparar arquiteturas, estratégias de amostragem e regras temporais antes da integração ao produto. Cada alteração relevante foi confrontada com uma referência controlada, reduzindo o risco de promover um modelo apenas por melhora pontual em uma única métrica.",
        101: "A arquitetura segue o modelo cliente-servidor mediado por API REST. O frontend React envia requisições ao backend FastAPI, responsável pela persistência, pelo processamento dos arquivos EDF e pela execução da IA. Esse desacoplamento permite manter a interface responsiva enquanto as tarefas computacionalmente intensivas são executadas no servidor.",
        105: "Na persistência, o PostgreSQL armazena metadados cadastrais e clínicos, enquanto arquivos EDF, artefatos dos modelos e mapas gerados são mantidos no sistema de arquivos e referenciados pelo banco. Não foram utilizados pgvector, banco vetorial ou formatos .npy/.h5 como persistência clínica. Essa delimitação reduz ambiguidades e corresponde à implementação do protótipo.",
        107: "A aplicação de IA na saúde enfrenta o problema da opacidade algorítmica. Em sistemas de apoio, uma métrica elevada não é suficiente para justificar confiança: também é necessário conhecer a origem dos dados, o protocolo de validação, as limitações do modelo e os elementos que influenciaram cada predição (AMANN et al., 2020; ADADI; BERRADA, 2018).",
        110: "Favorece a Interpretabilidade: o sistema associa os valores de relevância às features, aos canais e aos intervalos temporais da sequência analisada, permitindo confrontar a saída do modelo com o traçado exibido.",
        111: "Apoia a Revisão Clínica: a explicação não comprova uma justificativa fisiológica nem valida automaticamente a predição. Ela oferece uma indicação adicional para que o especialista examine se a região destacada é compatível com evento ictal, padrão interictal ou artefato.",
        112: "A combinação entre modelo temporal e XAI é, portanto, utilizada para tornar a saída computacional mais inspecionável, mantendo explícita a necessidade de supervisão humana e validação externa.",
        114: "A transformação digital em saúde depende de sistemas capazes de trocar e apresentar dados de forma consistente. Neste trabalho, a interoperabilidade foi delimitada ao suporte ao formato EDF e à comunicação por API REST. Integrações hospitalares baseadas em HL7 FHIR ou DICOM não foram implementadas e permanecem como possibilidade de evolução.",
        115: "A centralização web pode reduzir fragmentação operacional ao reunir cadastro, exame, visualização, análise e laudo em uma única interface. Entretanto, rapidez, segurança e benefício clínico não devem ser presumidos; essas propriedades precisam ser demonstradas por testes técnicos e estudos com usuários.",
        116: "Nesse contexto, a plataforma proposta se alinha à digitalização do fluxo de análise, mantendo o especialista responsável pela interpretação e pelo laudo. A automação fornece trechos suspeitos e explicações, mas não garante eliminação de erros nem maior precisão clínica sem validação apropriada.",
        122: "O desenvolvimento foi organizado nas etapas apresentadas na Figura 1, desde o levantamento de requisitos até a integração e a verificação técnica do protótipo.",
        124: "O levantamento de requisitos foi conduzido por revisão bibliográfica sobre análise computacional de EEG, análise documental do formato EDF, estudo do fluxo proposto no projeto anterior e refinamentos realizados durante a orientação acadêmica.",
        125: "Foram priorizados requisitos de ingestão, processamento de registros multicanais, visualização temporal, segurança, persistência e interpretabilidade. Não foram realizadas entrevistas estruturadas nem validação formal com neurologistas; por isso, os requisitos clínicos da versão atual são tratados como requisitos de um protótipo acadêmico.",
        126: "O benchmarking teve caráter bibliográfico e conceitual. Foram comparadas características descritas na literatura e em ferramentas de processamento de EEG, com atenção a visualização, seleção de canais, suporte a EDF e explicabilidade. Ferramentas que não tenham sido avaliadas de forma documentada não são apresentadas como objetos de comparação empírica.",
        127: "Como resultado, foram definidos os módulos de autenticação, pacientes, ingestão de EDF, visualização multicanal, análise preditiva, explicabilidade, configurações e emissão de laudo. A validação desses requisitos com profissionais da neurologia permanece como trabalho futuro.",
        129: "O desenvolvimento foi conduzido em ciclos incrementais individuais de implementação, teste e revisão. A orientação acadêmica apoiou a avaliação das entregas e o Git registrou a evolução do código. Nos ciclos de IA, hiperparâmetros e arquiteturas foram comparados por experimentos controlados antes da integração.",
        130: "O sistema utiliza arquitetura full-stack orientada a serviços, separando o cliente React do servidor FastAPI. Essa separação mantém responsabilidades claras e permite que o processamento matemático ocorra no backend sem transferir os arquivos brutos ou o modelo para o navegador.",
        132: "A implementação utilizou ferramentas com versões registradas para favorecer a reprodutibilidade do ambiente.",
        133: "O backend foi desenvolvido com Python 3.11.9, FastAPI 0.115.6 e SQLAlchemy 2.0.36. O PostgreSQL é destinado aos metadados, enquanto os EDF, modelos e mapas permanecem no sistema de arquivos. O ambiente Python é isolado por venv.",
        134: "O frontend foi desenvolvido com React 19.2.7, TypeScript, Vite 6.3.5 e Recharts 3.9.0. O Recharts foi adotado por sua integração declarativa com React; D3.js não integra a versão final do protótipo.",
        135: "O pipeline utiliza TensorFlow 2.16.2, MNE-Python 1.8.0, NumPy 1.26.4 e SHAP 0.46.0. O Git registra as alterações e o arquivo de dependências do Python, juntamente com o package-lock do frontend, apoia a reprodução do ambiente. Docker não foi adotado nesta versão.",
        139: "A implementação foi dividida em incrementos dependentes, mas com atividades parcialmente paralelas. Cada incremento recebeu critérios objetivos de aceite:",
        140: "Incremento I (Infraestrutura e Ingestão): concluído quando a API passou a autenticar usuários, cadastrar pacientes, receber, validar, armazenar e ler arquivos EDF, extraindo canais e taxa de amostragem.",
        141: "Incremento II (Núcleo de Inteligência): concluído quando o pipeline passou a gerar de forma reproduzível características tempo-frequência por canal, Entropia de Shannon e Dinâmica Simbólica, treinar a CNN-BiLSTM e produzir probabilidades temporais. A separação por paciente foi executada antes do janelamento e da amostragem para evitar vazamento de dados.",
        142: "Incremento III (Interface e Explicabilidade): concluído com a integração entre backend e frontend, incluindo visualização multicanal, seleção de canais, trechos suspeitos, sobreposição SHAP, configurações médicas e emissão de laudo.",
        143: "A cada ciclo foram executados testes automatizados e verificações sobre arquivos EDF reais. Alterações experimentais na IA somente foram promovidas quando apresentaram benefício no protocolo controlado de calibração.",
        147: "A verificação contemplou o comportamento do software e a avaliação científica do modelo, mantendo métricas de janela, EDF, evento e localização temporal em níveis distintos.",
        148: "Dados e métricas de IA: a base final foi a CHB-MIT, por oferecer registros contínuos multicanais em EDF e anotações de início e término das crises. Foram destinados ao treino 77 EDF de 17 grupos de pacientes e à calibração 13 EDF de chb16, chb19, chb20 e chb24. A separação ocorreu por paciente antes do janelamento. Foram medidos precisão, sensibilidade e F1 por EDF e evento, F1 localizado, IoU temporal e falsos alarmes por hora.",
        149: "Desempenho e carga: o processamento ponta a ponta foi verificado em EDF real de longa duração. O tempo aceitável para uso clínico não foi definido por profissionais da área; portanto, o trabalho não adota um limite arbitrário. Uma avaliação formal deve registrar hardware, duração do exame, tempo total, uso de memória e concorrência.",
        150: "Usabilidade e interpretabilidade: foram realizados testes técnicos da interface, do alinhamento temporal e da projeção do SHAP. Não foi conduzido estudo formal com médicos ou técnicos; essa ausência é apresentada como limitação e trabalho futuro.",
        153: "A versão avaliada foi implantada localmente como protótipo de pesquisa, com frontend, API, banco de dados e artefatos de IA executados no ambiente de desenvolvimento.",
        154: "A implantação em nuvem não foi validada. Embora a arquitetura permita execução em servidor, acesso remoto, escalabilidade, criptografia de transporte e operação multiusuário exigem configuração e avaliação próprias antes de qualquer uso externo.",
        155: "A API disponibiliza documentação Swagger/OpenAPI gerada pelo FastAPI. O repositório também registra variáveis de configuração, comandos de treinamento, calibração, avaliação e verificação dos artefatos.",
        156: "A avaliação utilizou registros públicos reais e desidentificados da CHB-MIT. Não foram coletados exames de hospital local nem recrutados participantes. Eventual estudo com dados identificáveis ou interação com pacientes e profissionais deverá ser submetido às instâncias éticas e institucionais aplicáveis.",
        182: "O software configura-se como um protótipo funcional de plataforma web de apoio à análise de EEG. Seu público-alvo potencial inclui neurologistas, neurofisiologistas e pesquisadores, mas a versão atual não foi validada clinicamente nem regularizada para uso assistencial. A proposta responde ao esforço envolvido na revisão de registros extensos e à necessidade de reunir visualização, análise computacional, explicabilidade e laudo em um fluxo único.",
        183: "A plataforma recebe arquivos EDF vinculados ao paciente, extrai metadados e apresenta os canais em uma linha temporal interativa. O pipeline CNN-BiLSTM produz probabilidades ictais por passo e agrega sequências sobrepostas em trechos suspeitos. O Gradient SHAP explica as features da sequência de pico e a interface projeta relevâncias positivas por canal e intervalo de segundos. Essa explicação não corresponde a uma atribuição por amostra bruta e deve ser interpretada como apoio à inspeção, não como justificativa fisiológica conclusiva.",
        185: "A engenharia de requisitos baseou-se em revisão bibliográfica, análise documental do EDF, experiência do projeto anterior e refinamento com a orientação acadêmica. Não houve entrevista estruturada nem validação formal com neurologistas. Assim, os requisitos clínicos representam hipóteses de projeto para um protótipo e deverão ser avaliados com profissionais antes de uma eventual evolução assistencial.",
        202: "3 - Segurança e Privacidade: o sistema deve restringir rotas clínicas a usuários autenticados, armazenar senhas por hash e evitar exposição desnecessária de dados. A versão atual não implementa consentimento, política automatizada de retenção, criptografia específica dos EDF em repouso ou trilha completa de auditoria; esses controles são necessários antes do uso com dados identificáveis.",
        208: "Ao solicitar a análise preditiva, o médico seleciona um exame previamente importado. A extração de métricas e a execução do modelo constituem operações internas obrigatórias do sistema, e não casos de uso iniciados diretamente pelo médico. A visualização da explicação permanece opcional durante a revisão do resultado.",
        307: "Para representar a persistência implementada no protótipo, o modelo relaciona quatro entidades principais: Usuario, Paciente, Exame e Predicao_IA. O texto do laudo permanece associado ao Exame, enquanto as saídas do modelo e a referência do SHAP são armazenadas em Predicao_IA. Não existe entidade específica de auditoria na versão atual.",
        308: "A entidade Usuario armazena dados profissionais e o hash da senha dos médicos autorizados. Paciente centraliza os dados cadastrais e se relaciona aos exames registrados. Embora o sistema possua autenticação, a implementação atual não equivale a conformidade integral com a LGPD.",
        309: "A entidade Exame registra o arquivo EDF, os metadados extraídos, o estado do processamento e o texto livre do laudo. Depois da emissão, o laudo torna-se imutável. Não há rascunho persistente, campos clínicos estruturados ou assinatura digital nesta versão.",
        310: "A entidade Predicao_IA registra o resultado analítico, o score, os trechos suspeitos e a referência do mapa SHAP. Arquivos EDF, modelos e imagens permanecem no sistema de arquivos e são referenciados pelo banco.",
        337: "A Figura 5 apresenta a interface de recuperação de senha. Na versão acadêmica, o fluxo é apenas simulado: a API devolve uma resposta neutra, mas não envia mensagem nem redefine credenciais. A recuperação operacional por e-mail permanece como trabalho futuro.",
        365: "A Figura 7 apresenta a interface de gestão de pacientes, dedicada ao registro e ao acompanhamento dos prontuários. A tela reúne pesquisa, cadastro, histórico de exames e situação do laudo, permitindo organizar o fluxo de trabalho do protótipo.",
        396: "A Figura 10 apresenta a tela de configurações do sistema, organizada em dados profissionais, preferências de análise e informações técnicas. O médico pode ajustar o limiar de confiança, selecionar a montagem padrão de canais e ativar ou desativar a exibição do SHAP. Esses controles personalizam a visualização, mas não substituem a calibração científica do modelo.",
        406: "O desenvolvimento foi organizado conforme o cronograma apresentado no Quadro 9, e os recursos empregados estão descritos no Quadro 8. O acompanhamento incremental permitiu revisar requisitos, implementar o sistema e executar experimentos de IA ao longo de 2026.",
        440: "META PLATFORMS. React: the library for web and native user interfaces. Disponível em: https://react.dev/. Acesso em: 13 out. 2026.",
        447: "REISMAN, M. EHRs: the challenge of making electronic data usable and interoperable. Pharmacy and Therapeutics, v. 42, n. 9, p. 572-575, 2017. Disponível em: https://pmc.ncbi.nlm.nih.gov/articles/PMC5565131/. Acesso em: 13 out. 2026.",
    }
    for index, text in replacements.items():
        set_text(p[index], text)

    # Renumber existing figures because the methodology diagram becomes Figure 1.
    for paragraph in doc.paragraphs:
        if paragraph is p[122]:
            continue
        has_drawing = bool(paragraph._p.xpath(".//w:drawing") or paragraph._p.xpath(".//w:pict"))
        if has_drawing:
            for run in paragraph.runs:
                original_text = run.text
                text = original_text
                for number in range(10, 0, -1):
                    text = re.sub(rf"\bFigura {number}\b", f"Figura {number + 1}", text)
                if text != original_text:
                    run.text = text
        else:
            text = paragraph.text
            for number in range(10, 0, -1):
                text = re.sub(rf"\bFigura {number}\b", f"Figura {number + 1}", text)
            if text != paragraph.text:
                set_text(paragraph, text)

    image_p = insert_after(p[122])
    image_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_p.add_run().add_picture(str(DIAGRAM), width=Inches(6.3))
    caption = insert_after(image_p, "Figura 1 - Etapas executadas no desenvolvimento e na avaliação do projeto.")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.line_spacing = 1.0
    source = insert_after(caption, "Fonte: autor, 2026.")
    source.alignment = WD_ALIGN_PARAGRAPH.CENTER
    source.paragraph_format.line_spacing = 1.0

    regulatory = insert_after(
        p[69],
        "No Brasil, a eventual disponibilização de software com finalidade médica deve considerar a "
        "regulamentação aplicável a Software como Dispositivo Médico, incluindo a RDC nº 657, de 24 de "
        "março de 2022, da Agência Nacional de Vigilância Sanitária. O protótipo desenvolvido neste trabalho "
        "não foi submetido à regularização sanitária. No âmbito da interface, as recomendações do eMAG e das "
        "WCAG constituem referências para evolução da acessibilidade, sem que se alegue conformidade formal "
        "antes de uma auditoria específica.",
    )
    regulatory.style = doc.styles["normal"]

    # Correct use-case tables according to the reviewers and the implemented behavior.
    predictive = doc.tables[4]
    predictive.cell(5, 1).text = "O arquivo EDF correspondente deve estar previamente importado e carregado no backend."
    predictive.cell(6, 1).text = "As probabilidades temporais, os trechos suspeitos e os metadados da análise são persistidos e disponibilizados na interface."

    xai = doc.tables[5]
    xai.cell(5, 1).text = "A análise preditiva descrita no Quadro 5 deve estar concluída e possuir uma sequência de pico válida."
    xai.cell(9, 1).text = "2 - Associar os valores positivos de Gradient SHAP às features, aos canais e aos intervalos temporais da sequência de pico."
    xai.cell(10, 1).text = "3 - Projetar as relevâncias por canal e intervalo de segundos sobre o traçado exibido, mantendo explícito que a explicação não é calculada por amostra bruta."

    report = doc.tables[6]
    report.cell(11, 1).text = "4 - Salvar o texto associado ao exame, alterar o status para Concluído e impedir novas edições."
    report.cell(12, 1).text = "1 - Impedir a emissão se o texto estiver vazio. 2 - Informar que não há rascunho persistente, campos estruturados ou assinatura digital nesta versão."

    for index in (293, 294, 295):
        if not p[index].text.startswith("- "):
            set_text(p[index], "- " + p[index].text)
        remove_numbering(p[index])

    # Apply readable academic formatting while preserving cover, headings, figures and tables.
    for index, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.strip()
        if index < 65 or not text:
            continue
        if paragraph.style.name.startswith("Heading"):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.paragraph_format.space_before = Pt(12)
            paragraph.paragraph_format.space_after = Pt(6)
            continue
        if text.startswith(("Figura ", "Fonte:", "Quadro ")):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.paragraph_format.first_line_indent = Cm(0)
            continue
        if index >= 415:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.paragraph_format.space_after = Pt(6)
            paragraph.paragraph_format.first_line_indent = Cm(0)
        else:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            paragraph.paragraph_format.line_spacing = 1.5
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            if not text.startswith(("Incremento ", "1 -", "2 -", "3 -", "4 -", "5 -", "6 -", "7 -", "8 -", "9 -")):
                paragraph.paragraph_format.first_line_indent = Cm(1.25)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                cell.vertical_alignment = 1
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.line_spacing = 1.0
                    paragraph.paragraph_format.space_after = Pt(2)
                    for run in paragraph.runs:
                        run.font.size = Pt(9.5)

    # Keep the first interface description and its anchored image away from the page boundary.
    p[325].paragraph_format.page_break_before = True

    # Insert authoritative references in alphabetical order.
    brasil_refs = [
        "BRASIL. Agência Nacional de Vigilância Sanitária. Resolução da Diretoria Colegiada - RDC nº 657, de 24 de março de 2022. Dispõe sobre a regularização de software como dispositivo médico. Brasília, DF: Anvisa, 2022. Disponível em: https://www.gov.br/anvisa/. Acesso em: 13 out. 2026.",
        "BRASIL. Agência Nacional de Vigilância Sanitária. Software como dispositivo médico: perguntas e respostas. Brasília, DF: Anvisa, 2022. Disponível em: https://www.gov.br/anvisa/pt-br/centraisdeconteudo/publicacoes/produtos-para-a-saude/manuais/software-como-dispositivo-medico-perguntas-e-respostas. Acesso em: 13 out. 2026.",
        "BRASIL. Governo Digital. Modelo de Acessibilidade em Governo Eletrônico - eMAG, versão 3.1. Brasília, DF: Governo Federal. Disponível em: https://www.gov.br/governodigitallogin/pt-br/acessibilidade-e-usuario/acessibilidade-digital/modelo-de-acessibilidade. Acesso em: 13 out. 2026.",
    ]
    for text in brasil_refs:
        reference = insert_before(p[424], text)
        reference.alignment = WD_ALIGN_PARAGRAPH.LEFT
        reference.paragraph_format.line_spacing = 1.0
        reference.paragraph_format.space_after = Pt(6)
        reference.paragraph_format.first_line_indent = Cm(0)

    shoeb = insert_before(
        p[451],
        "SHOEB, A. H. Application of machine learning to epileptic seizure onset detection and treatment. 2009. Tese (Doutorado) - Massachusetts Institute of Technology, Cambridge, 2009.",
    )
    shoeb.alignment = WD_ALIGN_PARAGRAPH.LEFT
    shoeb.paragraph_format.line_spacing = 1.0
    shoeb.paragraph_format.space_after = Pt(6)
    shoeb.paragraph_format.first_line_indent = Cm(0)

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
