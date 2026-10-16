from __future__ import annotations

import re
import shutil
import zipfile
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentType
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "tcc" / "TCC_II_Eduardo_Estilo_Karollyne_IFG_Final.docx"
TARGET = ROOT / "docs" / "tcc" / "overleaf_eeg_xai"
IMAGES = TARGET / "images"


def iter_blocks(parent: DocumentType):
    for child in parent.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, parent)
        elif child.tag == qn("w:tbl"):
            yield Table(child, parent)


def esc(value: str, *, trim: bool = True) -> str:
    leading = bool(value[:1].isspace())
    trailing = bool(value[-1:].isspace())
    value = value.replace("\u00a0", " ").replace("\v", " ")
    value = re.sub(r"\s+", " ", value).strip()
    replacements = {
        "\\": r"\textbackslash{}",
        "&": r"\&",
        "%": r"\%",
        "$": r"\$",
        "#": r"\#",
        "_": r"\_",
        "{": r"\{",
        "}": r"\}",
        "~": r"\textasciitilde{}",
        "^": r"\textasciicircum{}",
    }
    escaped = "".join(replacements.get(c, c) for c in value)
    if not trim and escaped:
        escaped = (" " if leading else "") + escaped + (" " if trailing else "")
    return escaped


def paragraph_text(paragraph: Paragraph) -> str:
    # python-docx versions before hyperlink support omit linked text from runs.
    if paragraph._p.xpath(".//w:hyperlink"):
        visible = "".join(node.text or "" for node in paragraph._p.xpath(".//w:t"))
        return esc(visible)
    parts: list[str] = []
    for run in paragraph.runs:
        text = esc(run.text, trim=False)
        if not text:
            continue
        if run.bold:
            text = rf"\textbf{{{text}}}"
        if run.italic:
            text = rf"\textit{{{text}}}"
        parts.append(text)
    return "".join(parts) or esc(paragraph.text)


def extract_images(document: DocumentType) -> dict[str, str]:
    IMAGES.mkdir(parents=True, exist_ok=True)
    mapping: dict[str, str] = {}
    for idx, rel in enumerate(
        (r for r in document.part.rels.values() if "image" in r.reltype), start=1
    ):
        suffix = Path(rel.target_ref).suffix.lower() or ".png"
        name = f"figura_{idx:02d}{suffix}"
        (IMAGES / name).write_bytes(rel.target_part.blob)
        mapping[rel.rId] = name
    return mapping


def paragraph_images(paragraph: Paragraph, image_map: dict[str, str]) -> list[tuple[str, float]]:
    result = []
    for blip in paragraph._p.xpath(".//a:blip"):
        rid = blip.get(qn("r:embed"))
        if rid in image_map:
            extents = paragraph._p.xpath(".//wp:extent")
            width_cm = float(extents[0].get("cx")) / 360000 if extents else 13.6
            result.append((image_map[rid], width_cm))
    return result


def table_latex(table: Table) -> str:
    rows = []
    ncols = max((len(row.cells) for row in table.rows), default=1)
    if ncols == 13:
        widths = [0.20] + [0.055] * 12
        font_size, tabcolsep = r"\scriptsize", "2pt"
    elif ncols == 6:
        widths = [0.05, 0.25, 0.08, 0.20, 0.15, 0.16]
        font_size, tabcolsep = r"\fontsize{9}{11}\selectfont", "2pt"
    else:
        widths = [0.92 / ncols] * ncols
        font_size, tabcolsep = r"\fontsize{10}{12}\selectfont", "3pt"
    spec = "|" + "|".join(
        rf">{{\raggedright\arraybackslash}}p{{{width:.3f}\textwidth}}" for width in widths
    ) + "|"
    rows.extend([r"\begingroup", font_size, rf"\setlength{{\tabcolsep}}{{{tabcolsep}}}", r"\renewcommand{\arraystretch}{1.15}"])
    # The caption preceding the longtable already advances the table counter.
    rows.append(r"\addtocounter{table}{-1}")
    rows.append(rf"\begin{{longtable}}{{{spec}}}")
    rows.append(r"\hline")
    for row_idx, row in enumerate(table.rows):
        values = []
        previous_tc = None
        for cell in row.cells:
            values.append("" if cell._tc is previous_tc else esc(cell.text))
            previous_tc = cell._tc
        values += [""] * (ncols - len(values))
        if ncols == 13 and row_idx == 0 and len(values) == 13:
            line = rf"\textbf{{{values[0]}}} & \multicolumn{{12}}{{c|}}{{\textbf{{{values[1]}}}}} \\ \hline"
        elif ncols == 6 and row_idx == len(table.rows) - 1 and values[0].startswith("TOTAL"):
            line = rf"\multicolumn{{6}}{{|r|}}{{\textbf{{{values[0]}}}}} \\ \hline"
        elif row_idx == 0:
            line = " & ".join(rf"\textbf{{{v}}}" for v in values) + r" \\ \hline"
        else:
            line = " & ".join(values) + r" \\ \hline"
        rows.append(line)
        if row_idx == 0:
            rows.extend([r"\endfirsthead", r"\hline", line, r"\endhead"])
    rows.append(r"\end{longtable}")
    rows.append(r"\endgroup")
    return "\n".join(rows)


def heading_command(style: str, text: str) -> str | None:
    clean = re.sub(r"^\s*\d+(?:\.\d+)*\s*", "", text).strip()
    if style == "Heading 1":
        if text.upper().startswith("REFER"):
            return r"\postextual\begin{SingleSpace}\chapter*{REFERÊNCIAS}\addcontentsline{toc}{chapter}{REFERÊNCIAS}"
        if text.upper().startswith("APÊNDICE"):
            return r"\end{SingleSpace}\begin{apendicesenv}\chapter{RECURSOS E CRONOGRAMA}"
        return rf"\chapter{{{esc(clean)}}}"
    if style == "Heading 2":
        clean = re.sub(r"^[A-Z]\.\d+\s*", "", clean).strip()
        return rf"\Needspace{{12\baselineskip}}\section{{{esc(clean)}}}"
    if style == "Heading 3":
        return rf"\subsection{{{esc(clean)}}}"
    return None


def build_body(document: DocumentType, image_map: dict[str, str]) -> str:
    out: list[str] = []
    seen_main = False
    pending_caption = ""
    list_open = False
    last_was_figure = False
    appendix_open = False
    references_open = False

    def close_list() -> None:
        nonlocal list_open
        if list_open:
            out.append(r"\end{itemize}")
            list_open = False

    for block in iter_blocks(document):
        if isinstance(block, Table):
            if not seen_main:
                continue
            close_list()
            out.append(table_latex(block))
            last_was_figure = False
            continue

        text = block.text.strip()
        images = paragraph_images(block, image_map)
        style = block.style.name if block.style else ""

        if style == "Heading 1" and text.startswith("1 "):
            seen_main = True
            out.extend([r"\cleardoublepage", r"\textual", r"\pagestyle{tcc}\aliaspagestyle{chapter}{tcc}"])

        if images:
            close_list()
            for image, width_cm in images:
                caption = re.sub(r"^Figura\s+\d+\s*[–-]\s*", "", pending_caption).strip()
                out.extend(
                    [
                        r"\begin{figure}[htbp]",
                        r"\centering",
                        rf"\caption{{{esc(caption or Path(image).stem)}}}",
                        rf"\includegraphics[width={width_cm:.2f}cm,height=0.72\textheight,keepaspectratio]{{images/{image}}}\par",
                        r"{\fontsize{10}{12}\selectfont Fonte: Elaborado pelo autor (2026).\par}",
                        r"\end{figure}",
                    ]
                )
            pending_caption = ""
            last_was_figure = True
            continue

        if not seen_main:
            continue

        if not text:
            close_list()
            out.append("")
            continue

        heading = heading_command(style, text)
        if heading:
            close_list()
            if text.upper().startswith("REFER"):
                references_open = True
            elif text.upper().startswith("APÊNDICE"):
                references_open = False
            out.append(heading)
            if text.upper().startswith("APÊNDICE"):
                appendix_open = True
            last_was_figure = False
            continue

        if re.match(r"^Figura\s+\d+\s*[–-]", text):
            pending_caption = text
            continue
        if text.startswith("Fonte:") and last_was_figure:
            last_was_figure = False
            continue
        if text.startswith("Quadro "):
            close_list()
            caption = re.sub(r"^Quadro\s+\d+\s*[–-]\s*", "", text).strip().rstrip(".")
            out.append(r"\FloatBarrier\Needspace{9\baselineskip}")
            out.append(rf"\captionof{{table}}{{{esc(caption)}}}\nopagebreak[4]")
            last_was_figure = False
            continue
        if text.startswith("•"):
            if not list_open:
                out.append(r"\begin{itemize}[leftmargin=1.25cm,itemsep=0pt,topsep=0pt]")
                list_open = True
            out.append(rf"\item {esc(text[1:].strip())}")
            continue
        close_list()
        if text.startswith("Fonte:"):
            out.append(rf"\nopagebreak[4]{{\centering\fontsize{{10}}{{12}}\selectfont {esc(text)}\par}}")
            last_was_figure = False
            continue
        if re.match(r"^\d+\s*[–-]\s*", text):
            out.append(rf"\noindent {paragraph_text(block)}\par")
            continue
        if references_open:
            visible = "".join(node.text or "" for node in block._p.xpath(".//w:t"))
            pieces = re.split(r"(https?://[^\s<>]+)", visible)
            formatted = "".join(
                rf"\url{{{piece}}}" if piece.startswith(("https://", "http://")) else esc(piece, trim=False)
                for piece in pieces
            )
            out.append(r"{\raggedright\setlength{\parindent}{0pt} " + formatted + r"\par}\smallskip" + "\n")
        else:
            out.append(paragraph_text(block) + "\n")
        last_was_figure = False
    close_list()
    if appendix_open:
        out.append(r"\end{apendicesenv}")
    return "\n".join(out)


def pretext(document: DocumentType) -> str:
    p = document.paragraphs
    resumo = esc(p[20].text)
    palavras = esc(p[21].text.removeprefix("Palavras-chave:").strip())
    abstract = esc(p[23].text)
    keywords = esc(p[24].text.removeprefix("Keywords:").strip())
    siglas = []
    for paragraph in p[28:65]:
        raw = paragraph.text.strip()
        if " – " in raw:
            key, value = raw.split(" – ", 1)
        elif " - " in raw:
            key, value = raw.split(" - ", 1)
        else:
            continue
        siglas.append(rf"\item[{esc(key)}] {esc(value)}")
    agradecimentos = "\n\n".join(esc(x.text) for x in p[14:19] if x.text.strip())
    return f"""\\begin{{capa}}
{{\\linespread{{1.15}}\\selectfont
\\begin{{center}}
INSTITUTO FEDERAL DE EDUCAÇÃO, CIÊNCIA E TECNOLOGIA DE GOIÁS -- CÂMPUS URUAÇU\\par
DEPARTAMENTO DE ÁREAS ACADÊMICAS\\par
CURSO DE TECNOLOGIA EM ANÁLISE E DESENVOLVIMENTO DE SISTEMAS\\par
\\vspace*{{110pt}}
EDUARDO HENRIQUE GUIMARÃES MARQUES\\par
\\vspace*{{130pt}}
\\textbf{{DESENVOLVIMENTO DE PLATAFORMA WEB PARA APOIO À ANÁLISE DE EEG COM INTELIGÊNCIA ARTIFICIAL EXPLICÁVEL}}\\par
\\vfill
URUAÇU/GO\\par
2026\\par
\\end{{center}}
}}
\\end{{capa}}

\\begin{{folhaderosto}}
{{\\linespread{{1.15}}\\selectfont
\\begin{{center}}
EDUARDO HENRIQUE GUIMARÃES MARQUES\\par
\\vspace*{{120pt}}
\\textbf{{DESENVOLVIMENTO DE PLATAFORMA WEB PARA APOIO À ANÁLISE DE EEG COM INTELIGÊNCIA ARTIFICIAL EXPLICÁVEL}}\\par
\\vspace*{{90pt}}
\\hspace*{{8.5cm}}\\begin{{minipage}}{{7.2cm}}
\\fontsize{{10}}{{12}}\\selectfont
\\begin{{SingleSpace}}
Trabalho de Conclusão de Curso apresentado ao curso de Tecnologia em Análise e Desenvolvimento de Sistemas do Departamento de Áreas Acadêmicas do Instituto Federal de Educação, Ciência e Tecnologia de Goiás -- Câmpus Uruaçu como requisito parcial para conclusão do curso.

\\vspace{{0.8cm}}
Orientador: Dr. Lynwood Livi de Souza
\\end{{SingleSpace}}
\\end{{minipage}}
\\par\\vfill
URUAÇU/GO\\par
2026\\par
\\end{{center}}
}}
\\end{{folhaderosto}}

\\begin{{agradecimentos}}
{agradecimentos}
\\end{{agradecimentos}}

\\begin{{resumo}}
{resumo}

\\vspace{{\\onelineskip}}
\\noindent\\textbf{{Palavras-chave}}: {palavras}
\\end{{resumo}}

\\begin{{resumo}}[ABSTRACT]
\\begin{{otherlanguage*}}{{english}}
{abstract}

\\vspace{{\\onelineskip}}
\\noindent\\textbf{{Keywords}}: {keywords}
\\end{{otherlanguage*}}
\\end{{resumo}}

\\pdfbookmark[0]{{\\listfigurename}}{{lof}}
\\listoffigures*
\\cleardoublepage
\\pdfbookmark[0]{{LISTA DE QUADROS}}{{loq}}
\\listoftables*
\\cleardoublepage

\\begin{{siglas}}
{chr(10).join(siglas)}
\\end{{siglas}}

\\pdfbookmark[0]{{\\contentsname}}{{toc}}
\\tableofcontents*
\\cleardoublepage
"""


def main() -> None:
    document = Document(SOURCE)
    TARGET.mkdir(parents=True, exist_ok=True)
    image_map = extract_images(document)
    (TARGET / "pretextuais.tex").write_text(pretext(document), encoding="utf-8")
    body = build_body(document, image_map)
    (TARGET / "conteudo.tex").write_text(body, encoding="utf-8")
    print(f"Exportados {len(image_map)} arquivos de imagem e {len(document.tables)} tabelas.")


if __name__ == "__main__":
    main()
