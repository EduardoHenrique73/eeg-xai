from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph


SOURCE = Path("docs/tcc/TCC II - Eduardo Marques - Versao Final Corrigida.docx")
OUTPUT = Path("docs/tcc/TCC II - Eduardo Marques - Versao Final ABNT.docx")

TITLE = "DESENVOLVIMENTO DE PLATAFORMA WEB PARA APOIO À ANÁLISE DE EEG COM INTELIGÊNCIA ARTIFICIAL EXPLICÁVEL"

RESUMO = (
    "A interpretação de exames de Eletroencefalograma (EEG) no diagnóstico da epilepsia exige a inspeção visual de longas séries temporais multicanais, um processo exaustivo e vulnerável à fadiga humana e à subjetividade entre avaliadores. Embora modelos de aprendizado profundo alcancem alta sensibilidade no reconhecimento de padrões bioelétricos, sua adoção clínica esbarra em duas barreiras críticas: a fragmentação entre ferramentas de visualização e a opacidade algorítmica das redes neurais, que operam como \"caixas-pretas\" sem justificativa inspecionável. Para superar essa lacuna, este trabalho desenvolveu e avaliou o protótipo funcional de uma plataforma web integrada (EEG-XAI) voltada ao apoio à triagem e à análise neurofisiológica. Estruturada sob arquitetura cliente-servidor mediada por API REST, a solução combina uma interface reativa em React, TypeScript e Recharts a um núcleo de processamento assíncrono em Python, FastAPI, MNE-Python, TensorFlow/Keras e PostgreSQL. No motor preditivo, implementou-se a arquitetura híbrida CNN-BiLSTM (versão v43) com fusão tardia de 23 canais, alimentada por sequências temporais de 15 características tempo-frequência por canal, incluindo entropia espectral. Para garantir transparência decisória, o método Gradient SHAP foi acoplado ao pipeline, projetando mapas de calor sobre o traçado original para destacar os canais e intervalos de segundos determinantes na sequência de pico. Nos experimentos com registros contínuos da base pública CHB-MIT, aplicando separação estrita por paciente antes do janelamento temporal para impedir vazamento de dados, o modelo atingiu na calibração F1 por arquivo EDF de 95,7%, F1 localizado de 90,9% e detecção de 15 em 20 eventos ictais, com 0,55 falso alarme por hora. Na avaliação histórica de desenvolvimento, obteve F1 por EDF de 82,4% e F1 localizado de 46,2%, evidenciando o desafio técnico da generalização interpaciente. Conclui-se que a plataforma unifica a ingestão de arquivos EDF, a navegação multicanal, a localização temporal de trechos suspeitos, a explicabilidade visual e a emissão de laudos, oferecendo uma ferramenta transparente de segunda opinião que apoia a revisão médica sem substituir o diagnóstico clínico."
)

ABSTRACT = (
    "Interpreting Electroencephalogram (EEG) exams for epilepsy diagnosis requires the visual inspection of extensive multichannel time series, an exhaustive process vulnerable to human fatigue and inter-observer subjectivity. Although deep learning models achieve high sensitivity in recognizing bioelectrical patterns, their clinical adoption faces two critical barriers: fragmentation among visualization tools and the algorithmic opacity of neural networks, which operate as \"black boxes\" without inspectable justification. To bridge this gap, this work developed and evaluated a functional integrated web platform prototype (EEG-XAI) designed to support neurophysiological triage and analysis. Built upon a REST API-mediated client-server architecture, the solution combines a reactive interface in React, TypeScript, and Recharts with an asynchronous processing core in Python, FastAPI, MNE-Python, TensorFlow/Keras, and PostgreSQL. At the predictive core, a hybrid CNN-BiLSTM architecture (version v43) with late fusion across 23 channels was implemented, fed by temporal sequences of 15 time-frequency features per channel, including spectral entropy. To ensure decision transparency, Gradient SHAP was integrated into the pipeline, projecting heatmaps onto the original trace to highlight the channels and time intervals driving the peak sequence. In experiments using continuous recordings from the public CHB-MIT dataset, with strict patient-level splitting before temporal windowing to prevent data leakage, the model achieved an EDF-level F1-score of 95.7%, a localized F1-score of 90.9%, and detection of 15 out of 20 ictal events in calibration, with 0.55 false alarms per hour. In the historical development evaluation, it reached an EDF-level F1-score of 82.4% and a localized F1-score of 46.2%, highlighting the challenge of inter-patient generalization. The platform integrates EDF ingestion, multichannel navigation, temporal localization of suspicious segments, visual explainability, and medical reporting, providing a transparent second-opinion tool that supports specialist review without replacing clinical diagnosis."
)

FIGURES = [
    "Etapas executadas no desenvolvimento e na avaliação do projeto",
    "Diagrama de Casos de Uso da Plataforma Web",
    "Fluxograma de Execução e Ciclo de Vida dos Dados da Plataforma",
    "Modelo Entidade Relacionamento",
    "Tela de login do sistema",
    "Janela modal de recuperação de senha do sistema",
    "Tela inicial (Dashboard) e visão geral do sistema",
    "Interface de gestão de pacientes e prontuários",
    "Janela modal de cadastro de novo paciente",
    "Interface central de visualização de EEG, análise de IA e laudo",
    "Interface de configurações do sistema e preferências de análise",
    "Painel principal da versão implementada",
    "Gestão de pacientes da versão implementada",
    "Configurações do perfil, montagem, limiar e SHAP",
    "Visualização multicanal com trecho suspeito e SHAP sobreposto por canal e tempo",
    "Área de emissão do parecer médico após a análise computacional",
    "Comparação das métricas da v43",
    "Eventos localizados e falsos alarmes por hora",
    "Verificação de integridade dos artefatos da v43",
]

TABLES = [
    "Documentação do caso de uso: Autenticar Usuário",
    "Documentação do caso de uso: Gerenciar Pacientes",
    "Documentação do caso de uso: Importar Exame (.EDF)",
    "Documentação do caso de uso: Visualizar Sinais de EEG",
    "Documentação do caso de uso: Solicitar Análise Preditiva",
    "Documentação do caso de uso: Visualizar Explicação (XAI)",
    "Documentação do caso de uso: Emitir Laudo Médico",
    "Resultados da v43 na calibração e na avaliação histórica",
    "Descrição Orçamentária",
    "Cronograma de trabalho dividido em meses e atividades a serem realizadas",
]

ABBREVIATIONS = [
    "ANVISA – Agência Nacional de Vigilância Sanitária",
    "API – Application Programming Interface (Interface de Programação de Aplicações)",
    "BiLSTM – Bidirectional Long Short-Term Memory (Memória de Longo e Curto Prazo Bidirecional)",
    "CDSS – Clinical Decision Support Systems (Sistemas de Suporte à Decisão Clínica)",
    "CHB-MIT – Children's Hospital Boston - Massachusetts Institute of Technology",
    "CNN – Convolutional Neural Network (Rede Neural Convolucional)",
    "CPF – Cadastro de Pessoas Físicas",
    "CRM – Conselho Regional de Medicina",
    "EDF – European Data Format",
    "EEG – Eletroencefalograma",
    "eMAG – Modelo de Acessibilidade em Governo Eletrônico",
    "FROC – Free-Response Receiver Operating Characteristic",
    "HTTP – Hypertext Transfer Protocol",
    "IA – Inteligência Artificial",
    "IFG – Instituto Federal de Educação, Ciência e Tecnologia de Goiás",
    "I/O – Input/Output (Entrada/Saída)",
    "IoU – Intersection over Union (Interseção sobre União)",
    "JSON – JavaScript Object Notation",
    "JWT – JSON Web Token",
    "LGPD – Lei Geral de Proteção de Dados Pessoais",
    "LSTM – Long Short-Term Memory (Memória de Longo e Curto Prazo)",
    "MER – Modelo Entidade-Relacionamento",
    "MLP – Multi-Layer Perceptron (Perceptron Multicamadas)",
    "P&D – Pesquisa e Desenvolvimento",
    "PIBITI – Programa Institucional de Bolsas de Iniciação em Desenvolvimento Tecnológico e Inovação",
    "RDC – Resolução da Diretoria Colegiada",
    "REST – Representational State Transfer",
    "RF – Requisito Funcional",
    "RNF – Requisito Não Funcional",
    "SADC – Sistema de Apoio à Decisão Clínica",
    "SHAP – Shapley Additive exPlanations",
    "SOA – Service-Oriented Architecture (Arquitetura Orientada a Serviços)",
    "SPA – Single Page Application",
    "UI – User Interface (Interface de Usuário)",
    "WCAG – Web Content Accessibility Guidelines",
    "WHO – World Health Organization (Organização Mundial da Saúde)",
    "XAI – Explainable Artificial Intelligence (Inteligência Artificial Explicável)",
]


def find_paragraph(document: Document, text: str) -> Paragraph:
    for paragraph in document.paragraphs:
        if paragraph.text.strip() == text:
            return paragraph
    raise RuntimeError(f"Parágrafo não encontrado: {text}")


def find_caption(document: Document, prefix: str, number: int) -> Paragraph:
    marker = f"{prefix} {number} "
    for paragraph in document.paragraphs:
        if paragraph.text.strip().startswith(marker):
            return paragraph
    raise RuntimeError(f"Legenda não encontrada: {prefix} {number}")


def move_before(anchor: Paragraph, element) -> None:
    anchor._p.addprevious(element)


def add_page_title(document: Document, anchor: Paragraph, text: str) -> Paragraph:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.page_break_before = True
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(18)
    run = paragraph.add_run(text)
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(12)
    move_before(anchor, paragraph._p)
    return paragraph


def add_body(document: Document, anchor: Paragraph, text: str, *, centered=False, single=False, indent=True) -> Paragraph:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if centered else WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.line_spacing = 1.0 if single else 1.5
    paragraph.paragraph_format.first_line_indent = Cm(0 if not indent else 1.25)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(12)
    move_before(anchor, paragraph._p)
    return paragraph


def add_blank(document: Document, anchor: Paragraph) -> None:
    paragraph = document.add_paragraph()
    move_before(anchor, paragraph._p)


def add_bookmark(paragraph: Paragraph, name: str, bookmark_id: int) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def set_pageref(paragraph: Paragraph, bookmark: str, cached: int) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), f" PAGEREF {bookmark} \\h ")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = str(cached)
    run.append(text)
    field.append(run)
    paragraph._p.append(field)


def style_list_table(table) -> None:
    table.autofit = False
    for row in table.rows:
        row.cells[0].width = Cm(15.3)
        row.cells[1].width = Cm(1.0)
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            borders = tc_pr.first_child_found_in("w:tcBorders")
            if borders is None:
                borders = OxmlElement("w:tcBorders")
                tc_pr.append(borders)
            for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
                tag = qn(f"w:{edge}")
                element = borders.find(tag)
                if element is None:
                    element = OxmlElement(f"w:{edge}")
                    borders.append(element)
                element.set(qn("w:val"), "nil")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.line_spacing = 1.0
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(11)


def add_list_page(document: Document, anchor: Paragraph, title: str, prefix: str, items: list[str], base_pages: list[int], start_id: int) -> None:
    add_page_title(document, anchor, title)
    table = document.add_table(rows=len(items), cols=2)
    for index, item in enumerate(items):
        caption = find_caption(document, prefix, index + 1)
        bookmark = f"{prefix.lower()}_{index + 1}"
        add_bookmark(caption, bookmark, start_id + index)
        table.cell(index, 0).text = f"{prefix} {index + 1} – {item}"
        set_pageref(table.cell(index, 1).paragraphs[0], bookmark, base_pages[index])
        table.cell(index, 1).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    style_list_table(table)
    move_before(anchor, table._tbl)


def replace_text(paragraph: Paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def main() -> None:
    document = Document(SOURCE)
    sumario = find_paragraph(document, "SUMÁRIO")

    add_page_title(document, sumario, "FOLHA DE APROVAÇÃO")
    add_body(document, sumario, "EDUARDO HENRIQUE GUIMARÃES MARQUES", centered=True, indent=False)
    add_blank(document, sumario)
    add_body(document, sumario, TITLE, centered=True, indent=False)
    add_blank(document, sumario)
    add_body(document, sumario, "Trabalho de Conclusão de Curso apresentado ao Curso de Tecnologia em Análise e Desenvolvimento de Sistemas do Departamento de Áreas Acadêmicas do Instituto Federal de Educação, Ciência e Tecnologia de Goiás – Câmpus Uruaçu como requisito parcial para conclusão do curso.", indent=False)
    add_blank(document, sumario)
    add_body(document, sumario, "Aprovado em: ____ de ______________ de 2026.", centered=True, indent=False)
    add_blank(document, sumario)
    banca = add_body(document, sumario, "BANCA EXAMINADORA", centered=True, indent=False)
    banca.runs[0].bold = True
    for member in [
        "__________________________________________________\nProf. Dr. Lynwood Livi de Souza (Orientador)\nInstituto Federal de Goiás (IFG) – Câmpus Uruaçu",
        "__________________________________________________\nProf.(a) Avaliador(a) 1\nInstituto Federal de Goiás (IFG) – Câmpus Uruaçu",
        "__________________________________________________\nProf.(a) Avaliador(a) 2\nInstituto Federal de Goiás (IFG) – Câmpus Uruaçu",
    ]:
        add_blank(document, sumario)
        add_body(document, sumario, member, centered=True, single=True, indent=False)

    add_page_title(document, sumario, "AGRADECIMENTOS")
    for paragraph in [
        "Agradeço primeiramente a Deus, pela saúde, força e perseverança concedidas ao longo de toda a minha caminhada acadêmica.",
        "Aos meus pais e à minha família, pelo apoio incondicional, compreensão e incentivo constante em cada etapa da minha formação.",
        "Ao meu orientador, Prof. Dr. Lynwood Livi de Souza, pela orientação criteriosa, paciência, disponibilidade e pelas valiosas contribuições técnicas e científicas que guiaram a evolução deste trabalho desde a iniciação tecnológica até a conclusão deste curso.",
        "Ao Instituto Federal de Educação, Ciência e Tecnologia de Goiás (IFG) – Câmpus Uruaçu, e a todos os professores do Curso de Tecnologia em Análise e Desenvolvimento de Sistemas, pelos ensinamentos transmitidos e pela dedicação à educação pública e de qualidade.",
        "Aos colegas de curso e amigos que compartilharam desafios, aprendizados e momentos importantes durante a graduação, e aos membros da banca examinadora, pela disponibilidade e pelas contribuições para o aperfeiçoamento deste trabalho.",
    ]:
        add_body(document, sumario, paragraph)

    add_page_title(document, sumario, "RESUMO")
    add_body(document, sumario, RESUMO, single=True, indent=False)
    keywords = add_body(document, sumario, "Palavras-chave: eletroencefalograma; epilepsia; aprendizado profundo; CNN-BiLSTM; inteligência artificial explicável; SHAP.", single=True, indent=False)
    keywords.paragraph_format.space_before = Pt(12)
    keywords.runs[0].bold = False

    add_page_title(document, sumario, "ABSTRACT")
    add_body(document, sumario, ABSTRACT, single=True, indent=False)
    keywords_en = add_body(document, sumario, "Keywords: electroencephalogram; epilepsy; deep learning; CNN-BiLSTM; explainable artificial intelligence; SHAP.", single=True, indent=False)
    keywords_en.paragraph_format.space_before = Pt(12)

    figure_pages = [19, 23, 30, 31, 32, 33, 34, 35, 36, 37, 38, 39, 40, 41, 42, 43, 45, 45, 48]
    add_list_page(document, sumario, "LISTA DE FIGURAS", "Figura", FIGURES, figure_pages, 100)
    table_pages = [24, 25, 26, 27, 28, 28, 29, 44, 55, 55]
    add_list_page(document, sumario, "LISTA DE QUADROS", "Quadro", TABLES, table_pages, 200)

    add_page_title(document, sumario, "LISTA DE ABREVIATURAS E SIGLAS")
    for item in ABBREVIATIONS:
        add_body(document, sumario, item, single=True, indent=False)

    chapter_two = find_paragraph(document, "2 REFERENCIAL TEÓRICO")
    heading = document.add_paragraph()
    heading.style = find_paragraph(document, "1.2 Justificativa").style
    heading.add_run("1.3 Organização do Trabalho")
    chapter_two._p.addprevious(heading._p)
    organization = document.add_paragraph()
    organization.style = find_paragraph(document, "Este trabalho encontra-se estruturado em sete capítulos.").style if any(p.text.strip() == "Este trabalho encontra-se estruturado em sete capítulos." for p in document.paragraphs) else document.styles["Normal"]
    organization.add_run(
        "Este trabalho encontra-se estruturado em sete capítulos. O Capítulo 1 introduz o contexto clínico e computacional da análise de EEG, definindo os objetivos geral e específicos, bem como a justificativa da pesquisa. O Capítulo 2 apresenta o referencial teórico, abordando os fundamentos neurofisiológicos, a engenharia de software orientada a serviços, a persistência de dados biomédicos, as arquiteturas híbridas de aprendizado profundo com explicabilidade (XAI) e a transformação digital no apoio diagnóstico. O Capítulo 3 descreve a metodologia adotada, detalhando o levantamento de requisitos, as tecnologias utilizadas, os incrementos de desenvolvimento, os protocolos de teste e a implementação. O Capítulo 4 documenta a engenharia do software, contemplando os requisitos funcionais e não funcionais, os casos de uso, a arquitetura cliente-servidor, o fluxograma operacional, o modelo de dados e os protótipos de baixa fidelidade. O Capítulo 5 apresenta os resultados alcançados, englobando a interface real implementada, o protocolo experimental na base CHB-MIT, o desempenho do modelo v43, a análise de estabilidade entre seeds e a verificação dos artefatos. O Capítulo 6 discute as implicações dos achados, as limitações técnicas e as ameaças à validade do estudo. Por fim, o Capítulo 7 sintetiza as conclusões do trabalho e aponta direções para pesquisas futuras."
    )
    organization.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    organization.paragraph_format.first_line_indent = Cm(1.25)
    organization.paragraph_format.line_spacing = 1.5
    chapter_two._p.addprevious(organization._p)

    caption_11 = find_paragraph(document, "Figura 11 – Interface de configurações do sistema e preferências de análise.")
    paragraphs = document.paragraphs
    caption_11_index = next(index for index, paragraph in enumerate(paragraphs) if paragraph._p is caption_11._p)
    source_fig_11 = next(
        paragraph
        for paragraph in paragraphs[caption_11_index + 1 :]
        if paragraph.text.strip() == "Fonte: Elaborado pelo autor (2026)."
    )
    closing = document.add_paragraph()
    closing.style = document.styles["Normal"]
    closing.add_run("A partir da concepção dessas interfaces de baixa fidelidade, estabeleceram-se as diretrizes visuais e funcionais necessárias para guiar a etapa de codificação do frontend em React.js e sua integração com a API FastAPI, cujas telas reais implementadas são apresentadas na seção de resultados.")
    closing.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    closing.paragraph_format.first_line_indent = Cm(1.25)
    closing.paragraph_format.line_spacing = 1.5
    source_fig_11._p.addnext(closing._p)

    figure_descriptions = {
        "Figura 12 – Painel principal da versão implementada.": "A Figura 12 apresenta o painel principal (Visão geral) da versão implementada após a autenticação médica. A interface exibe a identificação do profissional logado na barra lateral e organiza no topo os indicadores de pacientes cadastrados, exames pendentes e laudos emitidos, além de disponibilizar atalhos de acesso rápido para abertura de prontuários e ajuste das preferências de análise.",
        "Figura 13 – Gestão de pacientes da versão implementada.": "A Figura 13 ilustra a tela implementada para a gestão de pacientes. O módulo disponibiliza um campo de busca filtrada por nome ou CPF, o botão de ação para cadastro de novos registros e a listagem tabular dos pacientes ativos, integrando botões em cada linha para editar dados, excluir registros ou encaminhar diretamente o prontuário para a importação de um novo exame de EEG.",
        "Figura 14 – Configurações do perfil, montagem, limiar e SHAP.": "A Figura 14 apresenta a interface implementada para as configurações do sistema. Nela, o especialista gerencia seus dados de perfil (nome, CRM e e-mail), seleciona os canais preferenciais da montagem longitudinal bipolar como referência padrão para novas análises e controla os parâmetros de exibição da IA e XAI, podendo ajustar o controle deslizante (slider) do limiar (threshold) de confiança e habilitar ou desabilitar a exibição dos mapas SHAP.",
        "Figura 15 – Visualização multicanal com trecho suspeito e SHAP sobreposto por canal e tempo.": "A Figura 15 demonstra a tela central de visualização multicanal de EEG operando sobre um registro real de 23 canais. Na coluna esquerda, a interface reúne os dados cadastrais e o histórico do paciente, juntamente com a área de upload para arquivos no formato .edf. No painel central, o traçado das séries temporais é renderizado com controles superiores de navegação e zoom; nele, a faixa laranja delimita o trecho suspeito e a intensidade vermelha representa a contribuição positiva do Gradient SHAP. A projeção está associada às janelas de features da sequência de pico; portanto, não corresponde a uma atribuição para cada amostra bruta do sinal. Na coluna direita, na aba de Análise, são exibidos o seletor de canais, o status do motor de IA e o Score de Suspeita acompanhado da observação explícita de que a métrica atua como apoio, e não como diagnóstico.",
        "Figura 16 – Área de emissão do parecer médico após a análise computacional.": "Por fim, a Figura 16 apresenta a área de emissão do parecer médico na coluna lateral direita (aba Laudo), mantendo o traçado de EEG e a sobreposição do SHAP visíveis ao centro para inspeção simultânea. O componente disponibiliza um campo textual para que o especialista descreva a interpretação clínica do exame correlacionando a saída da IA com o quadro do paciente, finalizando o fluxo por meio do botão de emissão do laudo.",
    }
    for caption_text, description in figure_descriptions.items():
        caption = find_paragraph(document, caption_text)
        previous = caption._p.getprevious()
        if previous is not None:
            paragraph = Paragraph(previous, caption._parent)
            replace_text(paragraph, description)

    toc = next(
        table
        for table in document.tables
        if any("1 INTRODUÇÃO" in cell.text for row in table.rows for cell in row.cells)
    )
    row_template = deepcopy(toc.rows[4]._tr)
    toc.rows[5]._tr.addprevious(row_template)
    cells = toc.rows[5].cells
    cells[0].text = "1.3 Organização do Trabalho"
    cells[1].text = "14"

    bookmark_id = 300
    for row in toc.rows:
        label = row.cells[0].text.strip()
        target = next((p for p in document.paragraphs if p.text.strip().lower() == label.lower()), None)
        if target is None:
            continue
        bookmark = f"toc_{bookmark_id}"
        add_bookmark(target, bookmark, bookmark_id)
        cached_text = row.cells[1].text.strip()
        cached = int(cached_text) + 7 if cached_text.isdigit() and label != "1.3 Organização do Trabalho" else 14
        set_pageref(row.cells[1].paragraphs[0], bookmark, cached)
        row.cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        bookmark_id += 1

    settings = document.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")

    section_two = document.sections[1]
    pg_num_type = section_two._sectPr.find(qn("w:pgNumType"))
    if pg_num_type is None:
        pg_num_type = OxmlElement("w:pgNumType")
        section_two._sectPr.append(pg_num_type)
    pg_num_type.set(qn("w:start"), "11")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT.resolve())


if __name__ == "__main__":
    main()
