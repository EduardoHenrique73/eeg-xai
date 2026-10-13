"""Acrescenta interface final, resultados, discussao e conclusao ao TCC revisado."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "tcc" / "TCC I - Eduardo Marques - Revisado.docx"
OUTPUT = ROOT / "docs" / "tcc" / "TCC II - Eduardo Marques - Versao Final.docx"
INTERFACE = ROOT / "docs" / "tcc" / "assets" / "interface"
RESULTS = ROOT / "docs" / "tcc" / "assets" / "resultados"


def move_before(paragraph, anchor) -> None:
    anchor._p.addprevious(paragraph._p)


def add_paragraph_before(doc, anchor, text="", style=None):
    paragraph = doc.add_paragraph(text, style=style)
    move_before(paragraph, anchor)
    return paragraph


def add_picture_before(doc, anchor, path: Path, width: Cm):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(path), width=width)
    move_before(paragraph, anchor)
    return paragraph


def add_caption(doc, anchor, text: str):
    paragraph = add_paragraph_before(doc, anchor, text)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(0)
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(10)
    source = add_paragraph_before(doc, anchor, "Fonte: Elaborado pelo autor (2026).")
    source.alignment = WD_ALIGN_PARAGRAPH.CENTER
    source.paragraph_format.space_after = Pt(8)
    for run in source.runs:
        run.font.name = "Arial"
        run.font.size = Pt(10)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def add_results_table(doc, anchor) -> None:
    table = doc.add_table(rows=1, cols=3)
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:color"), "D9D9D9")
        borders.append(border)
    tbl_pr.append(borders)
    headers = ["Métrica", "Calibração", "Histórico"]
    for idx, text in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = text
        set_cell_shading(cell, "17324D")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.bold = True
            run.font.name = "Arial"
            run.font.size = Pt(10)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    rows = [
        ("Precisão por EDF", "100,0%", "87,5%"),
        ("Sensibilidade por EDF", "91,7%", "77,8%"),
        ("F1 por EDF", "95,7%", "82,4%"),
        ("F1 localizado", "90,9%", "46,2%"),
        ("Eventos detectados", "15 de 20", "4 de 9"),
        ("Recall por evento", "75,0%", "44,4%"),
        ("F1 por evento", "71,4%", "17,8%"),
        ("Falsos alarmes por hora", "0,55", "2,08"),
    ]
    for row_index, values in enumerate(rows, start=1):
        cells = table.add_row().cells
        for col, value in enumerate(values):
            cells[col].text = value
            cells[col].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cells[col].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if col == 0 else WD_ALIGN_PARAGRAPH.CENTER
            for run in cells[col].paragraphs[0].runs:
                run.font.name = "Arial"
                run.font.size = Pt(10)
        if row_index % 2 == 0:
            for cell in cells:
                set_cell_shading(cell, "EFF6FF")
    table.columns[0].width = Cm(8.2)
    table.columns[1].width = Cm(4.1)
    table.columns[2].width = Cm(4.1)
    anchor._p.addprevious(table._tbl)


def add_page_break_before(doc, anchor) -> None:
    paragraph = add_paragraph_before(doc, anchor)
    paragraph.add_run().add_break(WD_BREAK.PAGE)


def main() -> None:
    doc = Document(SOURCE)
    for section in doc.sections:
        pg_mar = section._sectPr.find(qn("w:pgMar"))
        if pg_mar is None:
            continue
        for attribute in ("top", "right", "bottom", "left", "header", "footer", "gutter"):
            key = qn(f"w:{attribute}")
            value = pg_mar.get(key)
            if value and "." in value:
                pg_mar.set(key, str(round(float(value))))
    anchor = next(p for p in doc.paragraphs if p.text.strip().startswith("5 RECURSOS E CRONOGRAMA"))
    references = next(p for p in doc.paragraphs if p.text.strip().startswith("6 REFER"))

    doc.paragraphs[50].text = (
        "Trabalho de Conclusão de Curso apresentado ao curso de Tecnologia em Análise e Desenvolvimento de "
        "Sistemas do Departamento de Áreas Acadêmicas do Instituto Federal de Educação, Ciência e Tecnologia "
        "de Goiás - Câmpus Uruaçu como requisito parcial para conclusão do curso."
    )

    wireframes = next(p for p in doc.paragraphs if p.text.strip().startswith("4.8"))
    wireframes.text = "4.8 Protótipos de baixa fidelidade"
    wireframes.style = "Heading 2"

    add_page_break_before(doc, anchor)
    add_paragraph_before(doc, anchor, "4.9 Interface implementada", "Heading 2")
    add_paragraph_before(
        doc,
        anchor,
        "As telas finais substituem os protótipos como evidência da implementação. A identidade visual utiliza azul e branco, navegação lateral persistente e organização orientada ao fluxo clínico. As capturas a seguir foram obtidas da aplicação local em execução, com dados de demonstração.",
    )
    add_picture_before(doc, anchor, INTERFACE / "dashboard_final.png", Cm(16.5))
    add_caption(doc, anchor, "Figura 11 - Painel principal da versão implementada.")
    add_picture_before(doc, anchor, INTERFACE / "pacientes_final.png", Cm(16.5))
    add_caption(doc, anchor, "Figura 12 - Gestão de pacientes da versão implementada.")
    add_picture_before(doc, anchor, INTERFACE / "configuracoes_final.png", Cm(16.5))
    add_caption(doc, anchor, "Figura 13 - Configurações do perfil, montagem, limiar e SHAP.")
    add_picture_before(doc, anchor, INTERFACE / "exame_shap_final.png", Cm(16.5))
    add_caption(doc, anchor, "Figura 14 - Visualização multicanal com trecho suspeito e SHAP sobreposto por canal e tempo.")
    add_paragraph_before(
        doc,
        anchor,
        "Na Figura 14, a faixa laranja delimita o trecho suspeito e a intensidade vermelha representa contribuição positiva do Gradient SHAP. A projeção está associada às janelas de features da sequência de pico; portanto, não corresponde a uma atribuição para cada amostra bruta do sinal.",
    )
    add_picture_before(doc, anchor, INTERFACE / "laudo_final.png", Cm(16.5))
    add_caption(doc, anchor, "Figura 15 - Área de emissão do parecer médico após a análise computacional.")

    add_page_break_before(doc, anchor)
    add_paragraph_before(doc, anchor, "5 RESULTADOS", "Heading 1")
    add_paragraph_before(doc, anchor, "5.1 Protocolo experimental", "Heading 2")
    add_paragraph_before(
        doc,
        anchor,
        "O modelo final selecionado foi a versão v43 com seed 42. A arquitetura CNN-BiLSTM utiliza fusão tardia de 23 canais, 15 características tempo-frequência por canal, janelas de 4 segundos, passo de 2 segundos e sequências de 8 passos. O treinamento utilizou 77 EDF de 17 grupos de pacientes. A calibração foi realizada em 13 EDF dos grupos chb16, chb19, chb20 e chb24, separados antes do janelamento.",
    )
    add_paragraph_before(
        doc,
        anchor,
        "O ponto operacional foi definido exclusivamente na calibração, com limiar de 0,60, duração mínima de 18 segundos e fechamento de lacunas de 8 segundos. A seleção do checkpoint priorizou desempenho por evento sob limite de falsos alarmes por hora. O conjunto chb09, chb15 e chb18 foi mantido separado no experimento final, mas já havia sido consultado durante o desenvolvimento; por isso, seus resultados são classificados como avaliação histórica, e não como teste final intocado.",
    )

    add_paragraph_before(doc, anchor, "5.2 Desempenho do modelo", "Heading 2")
    add_results_table(doc, anchor)
    caption = add_paragraph_before(doc, anchor, "Quadro 10 - Resultados da v43 na calibração e na avaliação histórica.")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    source = add_paragraph_before(doc, anchor, "Fonte: Elaborado pelo autor a partir dos artefatos v43 (2026).")
    source.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_picture_before(doc, anchor, RESULTS / "metricas_v43.png", Cm(16.5))
    add_caption(doc, anchor, "Figura 16 - Comparação das métricas da v43.")
    add_paragraph_before(
        doc,
        anchor,
        "Na calibração, a v43 atingiu F1 por EDF de 95,7%, F1 localizado de 90,9% e detectou 15 de 20 eventos, com 0,55 falso alarme por hora. A avaliação histórica apresentou F1 por EDF de 82,4%, F1 localizado de 46,2% e detecção de 4 de 9 eventos, com 2,08 falsos alarmes por hora. A diferença indica perda de generalização entre pacientes e impede interpretar a calibração como estimativa de desempenho clínico externo.",
    )
    add_picture_before(doc, anchor, RESULTS / "eventos_falsos_alarmes_v43.png", Cm(16.5))
    add_caption(doc, anchor, "Figura 17 - Eventos localizados e falsos alarmes por hora.")

    add_paragraph_before(doc, anchor, "5.3 Verificação do software e dos artefatos", "Heading 2")
    add_paragraph_before(
        doc,
        anchor,
        "A integração foi verificada no fluxo local completo: autenticação, cadastro do paciente, leitura de EDF, visualização multicanal, recuperação de exame analisado, apresentação do score, intervalo suspeito, sobreposição SHAP e emissão de laudo. O frontend foi compilado com TypeScript e Vite. Os testes focados do backend totalizaram 22 casos aprovados; a suíte completa apresentou bloqueio na inicialização do TensorFlow no ambiente utilizado e permanece como pendência de fechamento.",
    )
    add_picture_before(doc, anchor, RESULTS / "terminal_verificacao_v43.png", Cm(16.0))
    add_caption(doc, anchor, "Figura 18 - Verificação de integridade dos artefatos da v43.")

    add_page_break_before(doc, anchor)
    add_paragraph_before(doc, anchor, "6 DISCUSSÃO", "Heading 1")
    add_paragraph_before(
        doc,
        anchor,
        "Os resultados demonstram que o protótipo executa a proposta de apoio à análise temporal de EEG e apresenta evidências interpretáveis no mesmo contexto visual do sinal. A avaliação por EDF, isoladamente, mostrou-se insuficiente: mesmo quando o arquivo foi classificado corretamente, o sistema pôde deixar de localizar eventos específicos ou produzir alarmes fora do intervalo ictal. Por esse motivo, o estudo passou a registrar métricas por evento, localização temporal e falsos alarmes por hora.",
    )
    add_paragraph_before(
        doc,
        anchor,
        "A fusão tardia preservou a identidade dos canais por mais etapas da rede e foi superior às tentativas que agregavam precocemente a dimensão espacial. Entretanto, crises com morfologia, duração ou topografia pouco representadas no treino continuaram recebendo score baixo. O pós-processamento temporal melhorou continuidade e reduziu fragmentação, mas não recuperou eventos que o primeiro estágio não representou.",
    )
    add_paragraph_before(
        doc,
        anchor,
        "O SHAP aumenta a inspecionabilidade ao destacar quais canais e janelas de features contribuíram para a sequência de pico. Essa explicação não comprova causalidade fisiológica, não substitui a interpretação do traçado e pode refletir correlações aprendidas pelo modelo. Sua utilidade deverá ser avaliada com especialistas em estudo posterior.",
    )
    add_paragraph_before(doc, anchor, "6.1 Limitações e ameaças à validade", "Heading 2")
    limitations = [
        "não existe conjunto CHB-MIT local completamente intocado, pois todos os grupos disponíveis foram utilizados em treino, calibração ou desenvolvimento histórico;",
        "a calibração contém apenas quatro grupos de pacientes e vinte eventos, quantidade insuficiente para estimar generalização clínica;",
        "não foi realizada validação externa em outra instituição ou base temporal multicanal equivalente;",
        "não houve estudo formal de usabilidade com neurologistas ou técnicos de EEG;",
        "o sistema é um protótipo local, sem validação de desempenho concorrente, conformidade integral com a LGPD ou regularização sanitária;",
        "a suíte completa de testes precisa ser reexecutada em ambiente isolado após resolver o bloqueio de inicialização do TensorFlow.",
    ]
    for item in limitations:
        p = add_paragraph_before(doc, anchor, f"• {item}")
        p.paragraph_format.left_indent = Cm(1.25)

    add_page_break_before(doc, anchor)
    add_paragraph_before(doc, anchor, "7 CONCLUSÃO", "Heading 1")
    add_paragraph_before(
        doc,
        anchor,
        "Este trabalho desenvolveu um protótipo web de apoio à análise de EEG que integra autenticação, gestão de pacientes, ingestão de EDF, visualização multicanal, detecção temporal com CNN-BiLSTM, explicabilidade por Gradient SHAP e emissão de laudo. A implementação atende ao objetivo de reunir o fluxo técnico em uma única aplicação e mantém explícito que a saída da IA representa apoio à decisão, não diagnóstico autônomo.",
    )
    add_paragraph_before(
        doc,
        anchor,
        "A v43 apresentou desempenho elevado na calibração, mas a queda observada na avaliação histórica confirmou que a generalização interpaciente permanece o principal gargalo. Portanto, o protótipo pode ser apresentado como resultado de Engenharia de Software e experimento de IA explicável, mas não como sistema clinicamente validado.",
    )
    add_paragraph_before(
        doc,
        anchor,
        "Como trabalhos futuros, recomenda-se obter uma coorte externa mantida intocada, ampliar a diversidade de eventos ictais, avaliar arquiteturas espaciais e temporais em ablações controladas, calibrar o ponto operacional sem acesso ao teste, medir usabilidade com profissionais, completar controles de segurança e privacidade e avaliar requisitos regulatórios antes de qualquer uso assistencial.",
    )

    anchor.text = "8 RECURSOS E CRONOGRAMA"
    anchor.style = "Heading 1"
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text.startswith("5.1 Recursos"):
            paragraph.text = "8.1 Recursos"
            paragraph.style = "Heading 2"
        elif text.startswith("5.2 Cronograma"):
            paragraph.text = "8.2 Cronograma"
            paragraph.style = "Heading 2"
    references.text = "9 REFERÊNCIAS"
    references.style = "Heading 1"

    toc_lines = [
        "1 INTRODUÇÃO ........................................................................................................ 4",
        "    1.1 Objetivos ........................................................................................................ 5",
        "    1.2 Justificativa ..................................................................................................... 6",
        "2 REFERENCIAL TEÓRICO ........................................................................................ 7",
        "    2.1 Contexto neurofisiológico e desafios ................................................................ 7",
        "    2.2 Desenvolvimento ágil e arquitetura ................................................................... 8",
        "    2.3 Tecnologias full-stack e persistência ................................................................. 8",
        "    2.4 Inteligência artificial híbrida e explicabilidade ................................................... 9",
        "    2.5 Transformação digital no diagnóstico .............................................................. 11",
        "3 METODOLOGIA ...................................................................................................... 12",
        "    3.1 Levantamento de requisitos ............................................................................ 12",
        "    3.2 Desenvolvimento ............................................................................................. 13",
        "    3.3 Tecnologias e ferramentas ............................................................................... 13",
        "    3.4 Desenvolvimento incremental .......................................................................... 14",
        "    3.5 Testes .............................................................................................................. 15",
        "    3.6 Implementação ................................................................................................. 15",
        "4 SOFTWARE ............................................................................................................ 16",
        "    4.1 Descrição ......................................................................................................... 16",
        "    4.2 Levantamento de requisitos ............................................................................ 16",
        "    4.3 Requisitos funcionais e não funcionais ............................................................ 17",
        "    4.4 Casos de uso ................................................................................................... 18",
        "    4.5 Arquitetura do sistema ..................................................................................... 25",
        "    4.6 Fluxograma ...................................................................................................... 26",
        "    4.7 Modelo entidade relacionamento ..................................................................... 30",
        "    4.8 Protótipos de baixa fidelidade ......................................................................... 31",
        "    4.9 Interface implementada .................................................................................... 39",
        "5 RESULTADOS ......................................................................................................... 41",
        "    5.1 Protocolo experimental .................................................................................... 41",
        "    5.2 Desempenho do modelo ................................................................................... 41",
        "    5.3 Verificação do software e dos artefatos ........................................................... 43",
        "6 DISCUSSÃO ............................................................................................................. 44",
        "7 CONCLUSÃO ........................................................................................................... 45",
        "8 RECURSOS E CRONOGRAMA ................................................................................ 45",
        "9 REFERÊNCIAS ......................................................................................................... 46",
    ]
    toc = doc.paragraphs[63]
    toc.text = "\n".join(toc_lines)
    toc.style = doc.styles["Normal"]
    toc.alignment = WD_ALIGN_PARAGRAPH.LEFT
    toc.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    toc.paragraph_format.space_after = Pt(0)
    for run in toc.runs:
        run.font.name = "Arial"
        run.font.size = Pt(9)

    for instruction in doc.element.body.xpath('.//w:instrText[contains(text(), "TOC")]'):
        container = instruction
        while container is not None and container.tag != qn("w:sdt"):
            container = container.getparent()
        if container is not None:
            container.getparent().remove(container)

    for paragraph in doc.paragraphs:
        if paragraph.style and paragraph.style.name.startswith("Heading"):
            for run in paragraph.runs:
                run.font.name = "Arial"
                run.font.color.rgb = RGBColor(0, 0, 0)
        elif paragraph.text.strip():
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE

    settings = doc.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    core = doc.core_properties
    core.title = "Plataforma Web para Analise de EEG com Inteligencia Artificial Explicavel"
    core.subject = "TCC II - versao final do prototipo e resultados experimentais"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
